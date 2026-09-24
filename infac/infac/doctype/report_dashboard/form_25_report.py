import frappe
from frappe.utils import add_days, getdate, date_diff
from openpyxl import Workbook, utils
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from io import BytesIO
from datetime import datetime
from calendar import monthrange
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

status_map = {
    'Permission Request' :'PR',
    'On Duty':'OD',
    'Half Day':'HD',
    "Absent": "A",
    "Holiday": "HH",
    "Weekly Off": "WW",
    "Present": "P",
    "None" : "",
    "Leave Without Pay": "LOP",
    "Casual Leave": "CL",
    "Earned Leave": "EL",
    "Sick Leave": "SL",
    "Emergency -1": 'EML-1',
    "Emergency -2": 'EML-2',
    "Paternal Leave": 'PL',
    "Marriage Leave":'ML',
    "Medical Leave" : 'MDL',
    "Paternity Leave":'PTL',
    "Education Leave":'EL',
    "Maternity Leave":'MTL',
    "Covid -19": "COV-19",
    "Privilege Leave": "PVL",
    "Compensatory Off": "C-OFF",
    "BEREAVEMENT LEAVE":'BL'
}

@frappe.whitelist()
def download():
    """Download Attendance Report as Excel"""
    filename = 'Form-25'
    xlsx_file = make_xlsx()
    frappe.response['filename'] = filename + '.xlsx'
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary'

@frappe.whitelist()
def make_xlsx():
    """Build Attendance Excel"""
    filters = frappe.local.form_dict
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    factory_text = """Name & Address of the Factory
            INFAC INDIA PRIVATE LTD
            No 113, Ellaimman Koil Street
            Padappai Village, Kundrathur Taluk
            Kanchipuram Dist, Tamil Nadu
            PIN 601 301"""
    ws.merge_cells("A1:B6")
    fcell = ws["A1"]
    fcell.value = factory_text
    fcell.font = Font(size=11, bold=True)
    fcell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    ws["C1"] = "Form No 25"
    ws["C1"].font = Font(size=18, bold=True)
    ws["C1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 40 

    ws["C2"] = "(Prescribed under rule no 103 of the factories rules , 1950)"
    ws["C2"].font = Font(size=11, bold=True)
    ws["C2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 30

    ws["C3"] = "Attendance Register for the month of (Month - Year)"
    ws["C3"].font = Font(size=20, bold=True)
    ws["C3"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 50

    ws.column_dimensions["C"].width = 50

    columns = get_columns(filters)
    header_row = 7
    header_font = Font(bold=True)
    align_center = Alignment(horizontal="center", vertical="center")
    fill_weekend = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    for col_idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=col.split(":")[0])
        cell.font = header_font
        cell.alignment = align_center

    data = get_data(filters)
    for row_idx, row in enumerate(data, start=header_row + 1):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = align_center
            if value in ['WW', 'HH']:
                cell.fill = fill_weekend

    for i, col in enumerate(columns, start=1):
        ws.column_dimensions[utils.get_column_letter(i)].width = 15
    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    xlsx_file.seek(0)
    return xlsx_file

def get_columns(filters):
    columns = ["Employee ID","Employee Name","Employment Type","Department","Employee Category","DOJ","Status"]
    dates = get_dates(filters.from_date, filters.to_date)
    for d in dates:
        dt = getdate(d)
        columns.append(dt.strftime('%d/%b'))
    columns += ["Present","Half Day","On Duty","Permission","Absent","Weekoff","Holiday","Paid Leave","LOP","COFF"]
    return columns

# def get_data(filters):
#     data = []
#     employees = get_employees(filters)

#     for emp in employees:
#         row = [emp.name, emp.employee_name, emp.employment_type, emp.department, emp.employee_category, emp.date_of_joining, ""]
#         dates = get_dates(filters.from_date, filters.to_date)
#         totals = {
#             "P":0, "HD":0, "OD":0, "PR":0, "A":0,
#             "WW":0, "HH":0, "PL":0, "LOP":0, "COFF":0, "CL":0
#         }
#         for date in dates:
#             att = frappe.db.get_value(
#                 "Attendance",
#                 {"employee": emp.name, "attendance_date": date, "docstatus": ("!=", 2)},
#                 ["status","on_duty_marked","permission_request","leave_type"]
#             ) or ''
#             if att:
#                 status, on_duty, permission, leave_type = att[0], att[1], att[2], att[3]
#                 status_code = status_map.get(status, "")

#                 if on_duty: 
#                     hh = check_holiday(date, emp.name)
#                     if hh: 
#                         row.append(hh); totals[hh]+=1
#                     else: 
#                         row.append("OD"); totals["OD"]+=1

#                 elif permission:  
#                     hh = check_holiday(date, emp.name)
#                     if hh: 
#                         row.append(hh); totals[hh]+=1
#                     else: 
#                         row.append("P/P"); 
#                         totals["P"]+=1
#                         totals["PR"]+=1

#                 elif status in ["Present","Half Day","Absent"]:
#                     hh = check_holiday(date, emp.name)
#                     if hh: 
#                         row.append(hh); totals[hh]+=1
#                     else:
#                         row.append(status_code)
#                         totals[status_code]+=1

#                 elif leave_type:
#                     hh = check_holiday(date, emp.name)
#                     if hh:
#                         row.append(hh); totals[hh]+=1
#                     else:
#                         leave_status = status_map.get(leave_type, leave_type)
#                         row.append(leave_status)

#                         if leave_status == "C-OFF":
#                             totals["COFF"] += 1
#                         elif leave_status == "CL":
#                             totals["CL"] += 1
#                         elif leave_status == "LOP":
#                             totals["LOP"] += 1
#                         else:
#                             totals["PL"] += 1
#                 else:
#                     row.append("-")
#             else:
#                 hh = check_holiday(date, emp.name)
#                 if hh:
#                     row.append(hh); totals[hh]+=1
#                 else:
#                     row.append("-")
#         row += [
#             totals["P"], totals["HD"], totals["OD"], totals["PR"],
#             totals["A"], totals["WW"], totals["HH"], totals["PL"],
#             totals["LOP"], totals["COFF"], totals["CL"]
#         ]
#         data.append(row)

#     return data


def get_data(filters):
    data = []
    employees = get_employees(filters)

    for emp in employees:
        row = [emp.name, emp.employee_name, emp.employment_type, emp.department, emp.employee_category, emp.date_of_joining, ""]
        dates = get_dates(filters.from_date, filters.to_date)
        
        totals = {
            "P":0, "HD":0, "OD":0, "PR":0, "A":0,
            "WW":0, "HH":0, "PL":0, "LOP":0, "COFF":0, "CL":0
        }

        for date in dates:
            hh = check_holiday(date, emp.name)  # check holiday first
            att = frappe.db.get_value(
                "Attendance",
                {"employee": emp.name, "attendance_date": date, "docstatus": ("!=", 2)},
                ["status","on_duty_marked","permission_request","leave_type"]
            ) or ''

            if hh:
                row.append(hh)
                totals[hh] += 1
            elif att:
                status, on_duty, permission, leave_type = att[0], att[1], att[2], att[3]
                status_code = status_map.get(status, "")

                if on_duty:
                    row.append("OD"); totals["OD"] += 1
                elif permission:
                    row.append("P/P"); totals["P"] += 1; totals["PR"] += 1
                elif status == "Half Day":
                    if leave_type:
                        row.append("P/L"); totals["PL"] += 0.5; totals["HD"] += 0.5
                    else:
                        row.append("P/A"); totals["P"] += 0.5; totals["HD"] += 0.5
                elif status in ["Present", "Absent"]:
                    row.append(status_code); totals[status_code] += 1
                elif leave_type:
                    leave_status = status_map.get(leave_type, leave_type)
                    row.append(leave_status)
                    if leave_status == "C-OFF":
                        totals["COFF"] += 1
                    elif leave_status == "CL":
                        totals["CL"] += 1
                    elif leave_status == "LOP":
                        totals["LOP"] += 1
                    else:
                        totals["PL"] += 1
                else:
                    row.append("-")
            else:
                row.append("-")

        row += [
            totals["P"], totals["HD"], totals["OD"], totals["PR"],
            totals["A"], totals["WW"], totals["HH"], totals["PL"],
            totals["LOP"], totals["COFF"]
        ]
        data.append(row)

    return data

def get_dates(from_date, to_date):
    no_of_days = date_diff(add_days(to_date,1), from_date)
    return [add_days(from_date, i) for i in range(no_of_days)]

def get_employees(filters):
    conditions = ""
    if filters.employee:
        conditions += f" and employee='{filters.employee}'"
    if filters.employment_type: 
        conditions += f" and employment_type='{filters.employment_type}'"

    employees = frappe.db.sql(f"""select name, employee_name, department, employment_type, employee_category, date_of_joining
                                   from `tabEmployee` where status='Active' {conditions}""", as_dict=True)
    left_employees = frappe.db.sql(f"""select name, employee_name, department, employment_type, employee_category, date_of_joining
                                       from `tabEmployee` where status='Left' and relieving_date>='{filters.from_date}' {conditions}""", as_dict=True)
    employees.extend(left_employees)
    return employees

def check_holiday(date, emp):
    holiday_list = frappe.db.get_value('Company','Infac India Private Limited','default_holiday_list')
    holiday = frappe.db.sql("""select `tabHoliday`.holiday_date, `tabHoliday`.weekly_off from `tabHoliday List`
                               left join `tabHoliday` on `tabHoliday`.parent = `tabHoliday List`.name
                               where `tabHoliday List`.name=%s and holiday_date=%s""", (holiday_list, date), as_dict=True)
    if holiday:
        return "WW" if holiday[0].weekly_off else "HH"
    return None
