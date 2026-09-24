import frappe
from frappe.utils import getdate
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
import io
from datetime import datetime

@frappe.whitelist()
def download(from_date, to_date):

    if not from_date or not to_date:
        frappe.throw("From Date and To Date are mandatory")

    if getdate(from_date) > getdate(to_date):
        frappe.throw("From Date should not be greater than To Date")

    data = frappe.db.sql("""
        SELECT
            pr.employee_id AS employee_id,
            pr.employee_name,
            pr.department,
            emp.department_line,
            pr.workflow_state,
            pr.permission_date,
            pr.from_time,
            pr.to_time,
            pr.session,
            pr.shift
        FROM `tabPermission Request` pr
        LEFT JOIN `tabEmployee` emp
            ON emp.name = pr.employee_id
        WHERE pr.permission_date BETWEEN %s AND %s
        ORDER BY pr.permission_date ASC
    """, (from_date, to_date), as_dict=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Permission Report"

    headers = [
        "S.No",
        "Emp ID",
        "Emp Name",
        "Department",
        "Department Line",
        "Status",
        "Permission Date",
        "From Time",
        "To Time",
        "Session",
        "Shift"
    ]

    ws.append(headers)

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 25
    ws.column_dimensions["F"].width = 25
    ws.column_dimensions["G"].width = 20
    ws.column_dimensions["H"].width = 20
    ws.column_dimensions["I"].width = 20

    for idx, row in enumerate(data, start=1):

        excel_row = ws.max_row + 1
        raw_hours = row.get("hours")

        clean_time = None
        if raw_hours:
            if isinstance(raw_hours, str):
                clean_time = datetime.strptime(raw_hours.split(".")[0], "%H:%M:%S").time()
            else:
                clean_time = raw_hours.replace(microsecond=0)
        ws.append([
            idx,
            row.get("employee_id"),
            row.get("employee_name"),
            row.get("department"),
            row.get("department_line"),
            row.get("workflow_state"),
            row.get("permission_date"),
            # clean_time,
            row.get("from_time"),
            row.get("to_time"),
            row.get("session"),
            row.get("shift"),
        ])

        for col in range(1, 12):
            cell = ws.cell(row=excel_row, column=col)

            if col == 1:
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.alignment = Alignment(horizontal="left")

            cell.border = thin_border

        date_cell = ws.cell(row=excel_row, column=7)
        if row.get("permission_date"):
            date_cell.number_format = "DD-MM-YYYY"

        hour_cell = ws.cell(row=excel_row, column=8)
        if row.get("hours"):
            hour_cell.number_format = "HH:MM:SS"

    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)

    file_name = f"Permission_Report_{from_date}_to_{to_date}.xlsx"

    frappe.response["filename"] = file_name
    frappe.response["filecontent"] = file_stream.getvalue()
    frappe.response["type"] = "binary"