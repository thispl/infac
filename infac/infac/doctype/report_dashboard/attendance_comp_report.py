# from openpyxl import Workbook
# from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
# from datetime import datetime
# import frappe
# import io
# from frappe.utils import getdate, add_days

# @frappe.whitelist()
# def download(doc_date):

#     wb = Workbook()
#     ws = wb.active
#     ws.title = "Attendance Sheet"

#     thin_border = Border(
#         left=Side(style='thin'),
#         right=Side(style='thin'),
#         top=Side(style='thin'),
#         bottom=Side(style='thin')
#     )

#     total_columns = 23  # A to W

#     def generate_report_block(report_date, start_row):

#         date_obj = datetime.strptime(report_date, "%Y-%m-%d")
#         year_str = date_obj.year
#         month_str = date_obj.strftime("%b")
#         formatted_date = getdate(report_date).strftime("%d-%m-%Y")
        
#         # ---------------- DATE ROW (Above Header) ----------------
#         date_row = start_row - 1   # one row above header

#         ws.merge_cells(start_row=date_row, start_column=1, end_row=date_row, end_column=2)

#         ws.cell(row=date_row, column=1).value = formatted_date
#         ws.cell(row=date_row, column=1).font = Font(bold=True)
#         ws.cell(row=date_row, column=1).alignment = Alignment(horizontal="center", vertical="center")

#         yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

#         ws.cell(row=date_row, column=1).fill = yellow_fill
#         ws.cell(row=date_row, column=2).fill = yellow_fill

#         # Optional border
#         for col in range(1, 3):
#             ws.cell(row=date_row, column=col).border = thin_border

#         # ---------------- HEADER ----------------
#         ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=23)
#         ws.cell(row=start_row, column=1).value = f"Morning Attendance Report - {formatted_date}"
#         ws.cell(row=start_row, column=1).font = Font(bold=True)
#         ws.cell(row=start_row, column=1).alignment = Alignment(horizontal="center")

#         header_row1 = start_row + 1
#         header_row2 = start_row + 2

#         ws.merge_cells(start_row=header_row1, start_column=1, end_row=header_row2, end_column=3)
#         ws.cell(row=header_row1, column=1).value = "Department"
#         ws.cell(row=header_row1, column=1).alignment = Alignment(horizontal="center", vertical="center")

#         year_str = datetime.strptime(report_date, "%Y-%m-%d").year
#         month_str = datetime.strptime(report_date, "%Y-%m-%d").strftime("%b")

#         # Column 4 → Business Plan with Year above
#         ws.cell(row=header_row1, column=4).value = year_str
#         ws.cell(row=header_row2, column=4).value = "Business Plan"

#         # Column 5 → Available Count with Month above
#         ws.cell(row=header_row1, column=5).value = month_str
#         ws.cell(row=header_row2, column=5).value = "Available Count"

#         ws["D5"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
#         ws["E5"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

#         def create_group(title, col):
#             ws.merge_cells(start_row=header_row1, start_column=col,
#                            end_row=header_row1, end_column=col+2)
#             ws.cell(row=header_row1, column=col).value = title
#             ws.cell(row=header_row1, column=col).alignment = Alignment(horizontal="center")

#             ws.cell(row=header_row2, column=col).value = "Actual"
#             ws.cell(row=header_row2, column=col+1).value = "Present"
#             ws.cell(row=header_row2, column=col+2).value = "Gap"

#         col_pointer = 6
#         for shift in ["A", "G", "B", "C", "Available", "Staff"]:
#             create_group(shift, col_pointer)
#             col_pointer += 3

#         # ---------------- DATA ----------------
#         deps = frappe.get_all("Department",
#                               filters={"disabled": 0},
#                               fields=["name", "separate"],
#                               order_by="name")

#         normal_deps = [d for d in deps if not d.separate]
#         separate_deps = [d for d in deps if d.separate]

#         row_no = header_row2 + 1
#         start_data_row = row_no

#         def write_department(dep_name):

#             nonlocal row_no

#             ws.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=3)
#             ws.cell(row=row_no, column=1).value = dep_name
#             ws.cell(row=row_no, column=4).value = 0

#             available = frappe.db.count(
#                 "Employee",
#                 {
#                     "department": dep_name,
#                     "status": "Active",
#                     "employment_type": ["!=", "STAFF"]
#                 }
#             )
#             ws.cell(row=row_no, column=5).value = available

#             shifts = ["A", "G", "B", "C"]
#             col_pointer = 6

#             total_actual = total_present = total_gap = 0

#             for shift in shifts:

#                 shift_ass = frappe.db.sql("""
#                     SELECT COUNT(sa.employee)
#                     FROM `tabShift Assignment` sa
#                     INNER JOIN `tabEmployee` e ON e.name = sa.employee
#                     WHERE sa.docstatus = 1
#                     AND sa.shift_type = %s
#                     AND sa.start_date <= %s
#                     AND sa.end_date >= %s
#                     AND e.department = %s
#                     AND e.employment_type != 'STAFF'
#                 """, (shift, report_date, report_date, dep_name))[0][0]

#                 shift_pres = frappe.db.sql("""
#                     SELECT COUNT(DISTINCT ec.employee)
#                     FROM `tabEmployee Checkin` ec
#                     INNER JOIN `tabEmployee` e ON e.name = ec.employee
#                     WHERE DATE(ec.time) = %s
#                     AND e.department = %s
#                     AND e.employment_type != 'STAFF'
#                     AND ec.shift = %s
#                     AND ec.log_type = 'IN'
#                     AND ec.shift IS NOT NULL
#                     AND ec.shift != ''
#                 """, (report_date, dep_name, shift))[0][0]

#                 gap = shift_ass - shift_pres

#                 ws.cell(row=row_no, column=col_pointer).value = shift_ass
#                 ws.cell(row=row_no, column=col_pointer+1).value = shift_pres
#                 ws.cell(row=row_no, column=col_pointer+2).value = gap

#                 total_actual += shift_ass
#                 total_present += shift_pres
#                 total_gap += gap

#                 col_pointer += 3

#             # Shift total
#             ws.cell(row=row_no, column=col_pointer).value = total_actual
#             ws.cell(row=row_no, column=col_pointer+1).value = total_present
#             ws.cell(row=row_no, column=col_pointer+2).value = total_gap
#             col_pointer += 3

#             # Staff section
#             staff_available = frappe.db.count(
#                 "Employee",
#                 {
#                     "department": dep_name,
#                     "status": "Active",
#                     "employment_type": "STAFF"
#                 }
#             )

#             staff_present = frappe.db.sql("""
#                 SELECT COUNT(DISTINCT ec.employee)
#                 FROM `tabEmployee Checkin` ec
#                 INNER JOIN `tabEmployee` e ON e.name = ec.employee
#                 WHERE DATE(ec.time) = %s
#                 AND e.department = %s
#                 AND e.status = 'Active'
#                 AND e.employment_type = 'Staff'
#                 AND ec.log_type = 'IN'
#                 AND ec.shift IS NOT NULL
#                 AND ec.shift != ''
#             """, (report_date, dep_name))[0][0]

#             ws.cell(row=row_no, column=col_pointer).value = staff_available
#             ws.cell(row=row_no, column=col_pointer+1).value = staff_present
#             ws.cell(row=row_no, column=col_pointer+2).value = staff_available - staff_present

#             row_no += 1

#         # Write departments
#         for d in normal_deps:
#             write_department(d["name"])

#         normal_total_row = row_no
#         row_no += 1

#         for d in separate_deps:
#             write_department(d["name"])

#         separate_total_row = row_no
#         row_no += 1

#         overall_total_row = row_no

#         # Write row labels
#         ws.merge_cells(start_row=normal_total_row, start_column=1,
#                     end_row=normal_total_row, end_column=3)
#         ws.cell(normal_total_row, 1).value = "Total"

#         ws.merge_cells(start_row=separate_total_row, start_column=1,
#                     end_row=separate_total_row, end_column=3)
#         ws.cell(separate_total_row, 1).value = "Total"

#         ws.merge_cells(start_row=overall_total_row, start_column=1,
#                     end_row=overall_total_row, end_column=3)
#         ws.cell(overall_total_row, 1).value = "Overall Total"

#         # ---------------- TOTAL CALCULATION ----------------
#         for col in range(4, total_columns + 1):

#             normal_sum = sum(ws.cell(r, col).value or 0
#                              for r in range(start_data_row, normal_total_row))

#             separate_sum = sum(ws.cell(r, col).value or 0
#                                for r in range(normal_total_row+1, separate_total_row))

#             ws.cell(normal_total_row, col).value = normal_sum
#             ws.cell(separate_total_row, col).value = separate_sum
#             ws.cell(overall_total_row, col).value = normal_sum + separate_sum

#         # Border
#         for r in range(start_row, overall_total_row + 1):
#             for c in range(1, total_columns + 1):
#                 ws.cell(r, c).border = thin_border
#                 ws.cell(r, c).alignment = Alignment(horizontal="center", vertical="center")

#         return overall_total_row + 3

#     # First table
#     next_row = generate_report_block(doc_date, 2)

#     # Second table (doc_date - 1)
#     previous_date = add_days(doc_date, -1)
#     generate_report_block(previous_date, next_row)

#     file_stream = io.BytesIO()
#     wb.save(file_stream)
#     file_stream.seek(0)

#     frappe.response["filename"] = "Attendance_Comparison_Sheet.xlsx"
#     frappe.response["filecontent"] = file_stream.getvalue()
#     frappe.response["type"] = "binary"


from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from datetime import datetime
import frappe
import io
from frappe.utils import getdate, add_days

@frappe.whitelist()
def download(doc_date):

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Sheet"

    # ================= EXISTING BORDER =================
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # ================= NEW: THICK BORDER =================
    medium_border = Border(
        left=Side(style='medium'),
        right=Side(style='medium'),
        top=Side(style='medium'),
        bottom=Side(style='medium')
    )

    # ================= NEW COLORS =================
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    subheader_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    total_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    staff_fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")

    total_columns = 23  # A to W

    def generate_report_block(report_date, start_row):

        date_obj = datetime.strptime(report_date, "%Y-%m-%d")
        year_str = date_obj.year
        month_str = date_obj.strftime("%b")
        formatted_date = getdate(report_date).strftime("%d-%m-%Y")

        # ---------------- DATE ROW ----------------
        date_row = start_row - 1

        ws.merge_cells(start_row=date_row, start_column=1, end_row=date_row, end_column=2)
        ws.cell(row=date_row, column=1).value = formatted_date
        ws.cell(row=date_row, column=1).font = Font(bold=True)
        ws.cell(row=date_row, column=1).alignment = Alignment(horizontal="center", vertical="center")

        yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        ws.cell(row=date_row, column=1).fill = yellow_fill
        ws.cell(row=date_row, column=2).fill = yellow_fill

        for col in range(1, 3):
            ws.cell(row=date_row, column=col).border = thin_border

        # ---------------- TITLE ----------------
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=23)
        ws.cell(row=start_row, column=1).value = f"Morning Attendance Report - {formatted_date}"
        ws.cell(row=start_row, column=1).font = Font(bold=True)
        ws.cell(row=start_row, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=start_row, column=1).fill = header_fill  # NEW

        header_row1 = start_row + 1
        header_row2 = start_row + 2

        ws.merge_cells(start_row=header_row1, start_column=1, end_row=header_row2, end_column=3)
        ws.cell(row=header_row1, column=1).value = "Department"
        ws.cell(row=header_row1, column=1).alignment = Alignment(horizontal="center", vertical="center")

        ws.cell(row=header_row1, column=4).value = year_str
        ws.cell(row=header_row2, column=4).value = "Business Plan"

        ws.cell(row=header_row1, column=5).value = month_str
        ws.cell(row=header_row2, column=5).value = "Available Count"

        ws["D5"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws["E5"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        def create_group(title, col):
            ws.merge_cells(start_row=header_row1, start_column=col,
                           end_row=header_row1, end_column=col+2)
            ws.cell(row=header_row1, column=col).value = title
            ws.cell(row=header_row1, column=col).alignment = Alignment(horizontal="center")

            ws.cell(row=header_row2, column=col).value = "Actual"
            ws.cell(row=header_row2, column=col+1).value = "Present"
            ws.cell(row=header_row2, column=col+2).value = "Gap"

        col_pointer = 6
        for shift in ["A", "G", "B", "C", "Available", "Staff"]:
            create_group(shift, col_pointer)
            col_pointer += 3

        # ===== HEADER COLORS =====
        for c in range(1, total_columns + 1):
            ws.cell(header_row1, c).fill = header_fill
            ws.cell(header_row2, c).fill = subheader_fill
            ws.cell(header_row1, c).font = Font(bold=True)
            ws.cell(header_row2, c).font = Font(bold=True)

        # ---------------- DATA ----------------
        deps = frappe.get_all(
            "Department",
            filters={"disabled": 0, "hr_dashboard": 1},
            fields=["name", "separate"],
            order_by="name"
        )

        normal_deps = [d for d in deps if not d.separate]
        separate_deps = [d for d in deps if d.separate]

        row_no = header_row2 + 1
        start_data_row = row_no

        def write_department(dep_name):
            nonlocal row_no

            ws.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=3)
            ws.cell(row=row_no, column=1).value = dep_name
            # ===== FIX: merged department full border =====
            # for c in range(1, 4):  # only merged area
            #     ws.cell(row=row_no, column=c).border = medium_border

            available = frappe.db.count(
                "Employee",
                {
                    "department": dep_name,
                    "status": "Active",
                    "employment_type": ["!=", "STAFF"]
                }
            )
            ws.cell(row=row_no, column=5).value = available

            shifts = ["A", "G", "B", "C"]
            col_pointer = 6

            total_actual = total_present = total_gap = 0

            for shift in shifts:
                shift_ass = frappe.db.sql("""
                    SELECT COUNT(sa.employee)
                    FROM `tabShift Assignment` sa
                    INNER JOIN `tabEmployee` e ON e.name = sa.employee
                    WHERE sa.docstatus = 1
                    AND sa.shift_type = %s
                    AND sa.start_date <= %s
                    AND sa.end_date >= %s
                    AND e.department = %s
                    AND e.employment_type != 'STAFF'
                """, (shift, report_date, report_date, dep_name))[0][0]

                shift_pres = frappe.db.sql("""
                    SELECT COUNT(DISTINCT ec.employee)
                    FROM `tabEmployee Checkin` ec
                    INNER JOIN `tabEmployee` e ON e.name = ec.employee
                    WHERE DATE(ec.time) = %s
                    AND e.department = %s
                    AND e.employment_type != 'STAFF'
                    AND ec.shift = %s
                    AND ec.log_type = 'IN'
                    AND ec.shift IS NOT NULL
                    AND ec.shift != ''
                """, (report_date, dep_name, shift))[0][0]

                gap = shift_ass - shift_pres

                ws.cell(row=row_no, column=col_pointer).value = shift_ass
                ws.cell(row=row_no, column=col_pointer+1).value = shift_pres
                ws.cell(row=row_no, column=col_pointer+2).value = gap

                total_actual += shift_ass
                total_present += shift_pres
                total_gap += gap

                col_pointer += 3

            ws.cell(row=row_no, column=col_pointer).value = total_actual
            ws.cell(row=row_no, column=col_pointer+1).value = total_present
            ws.cell(row=row_no, column=col_pointer+2).value = total_gap
            col_pointer += 3

            staff_available = frappe.db.count(
                "Employee",
                {
                    "department": dep_name,
                    "status": "Active",
                    "employment_type": "STAFF"
                }
            )

            staff_present = frappe.db.sql("""
                SELECT COUNT(DISTINCT ec.employee)
                FROM `tabEmployee Checkin` ec
                INNER JOIN `tabEmployee` e ON e.name = ec.employee
                WHERE DATE(ec.time) = %s
                AND e.department = %s
                AND e.status = 'Active'
                AND e.employment_type = 'Staff'
                AND ec.log_type = 'IN'
                AND ec.shift IS NOT NULL
                AND ec.shift != ''
            """, (report_date, dep_name))[0][0]

            ws.cell(row=row_no, column=col_pointer).value = staff_available
            ws.cell(row=row_no, column=col_pointer+1).value = staff_present
            ws.cell(row=row_no, column=col_pointer+2).value = staff_available - staff_present
            available_start = col_pointer - 3


            # ===== STAFF COLOR =====
            for c in range(col_pointer, col_pointer + 3):
                ws.cell(row=row_no, column=c).fill = staff_fill
            for c in range(available_start, available_start + 3):
                ws.cell(row=row_no, column=c).fill = staff_fill

            row_no += 1

        for d in normal_deps:
            write_department(d["name"])

        normal_total_row = row_no
        row_no += 1

        for d in separate_deps:
            write_department(d["name"])

        separate_total_row = row_no
        row_no += 1

        overall_total_row = row_no

        ws.merge_cells(start_row=normal_total_row, start_column=1,
                       end_row=normal_total_row, end_column=3)
        ws.cell(normal_total_row, 1).value = "Total"

        ws.merge_cells(start_row=separate_total_row, start_column=1,
                       end_row=separate_total_row, end_column=3)
        ws.cell(separate_total_row, 1).value = "Total"

        ws.merge_cells(start_row=overall_total_row, start_column=1,
                       end_row=overall_total_row, end_column=3)
        ws.cell(overall_total_row, 1).value = "Overall Total"

        for col in range(4, total_columns + 1):
            normal_sum = sum(ws.cell(r, col).value or 0
                             for r in range(start_data_row, normal_total_row))
            separate_sum = sum(ws.cell(r, col).value or 0
                               for r in range(normal_total_row+1, separate_total_row))

            ws.cell(normal_total_row, col).value = normal_sum
            ws.cell(separate_total_row, col).value = separate_sum
            ws.cell(overall_total_row, col).value = normal_sum + separate_sum

        # ===== TOTAL ROW COLOR =====
        for c in range(1, total_columns + 1):
            for r in [normal_total_row, separate_total_row, overall_total_row]:
                ws.cell(r, c).fill = total_fill
                ws.cell(r, c).font = Font(bold=True)

        # ===== THICK OUTER BORDER =====
        # for r in range(start_row, overall_total_row + 1):
        #     for c in range(1, total_columns + 1):

        #         if (
        #             r == start_row or
        #             r == overall_total_row or
        #             c == 1 or
        #             c == total_columns
        #         ):
        #             ws.cell(r, c).border = medium_border
        #         else:
        #             ws.cell(r, c).border = thin_border

        #         ws.cell(r, c).alignment = Alignment(horizontal="center", vertical="center")
        # ===== CLEAN GROUP BORDER =====
    # ===== CLEAN GROUP BORDER (FINAL FIX) =====
        group_left_edges = [6, 9, 12, 15, 18, 21]
        group_right_edges = [8, 11, 14, 17, 20, 23]

        for r in range(start_row, overall_total_row + 1):
            for c in range(1, total_columns + 1):

                border_to_apply = thin_border

                # ✅ DEPARTMENT MERGED BOX (MOST IMPORTANT)
                if r >= start_data_row and r < overall_total_row and c in (1, 2, 3):
                    border_to_apply = medium_border

                # ===== OUTER BORDER =====
                elif (
                    r == start_row or
                    r == overall_total_row or
                    c == 1 or
                    c == total_columns
                ):
                    border_to_apply = medium_border

                # ===== GROUP BOX =====
                elif c in group_left_edges:
                    border_to_apply = Border(
                        left=Side(style='medium'),
                        right=Side(style='thin'),
                        top=Side(style='thin'),
                        bottom=Side(style='thin')
                    )
                elif c in group_right_edges:
                    border_to_apply = Border(
                        left=Side(style='thin'),
                        right=Side(style='medium'),
                        top=Side(style='thin'),
                        bottom=Side(style='thin')
                    )

                ws.cell(r, c).border = border_to_apply
                ws.cell(r, c).alignment = Alignment(horizontal="center", vertical="center")
        
        return overall_total_row + 3

    next_row = generate_report_block(doc_date, 2)
    previous_date = add_days(doc_date, -1)
    generate_report_block(previous_date, next_row)

    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)

    frappe.response["filename"] = "Attendance_Comparison_Sheet.xlsx"
    frappe.response["filecontent"] = file_stream.getvalue()
    frappe.response["type"] = "binary"