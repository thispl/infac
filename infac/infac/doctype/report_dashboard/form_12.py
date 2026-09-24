import frappe
from frappe.model.document import Document
from openpyxl import Workbook, utils
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.drawing.image import Image as xlImage
from io import BytesIO
import requests
from frappe.utils import add_days, getdate
from urllib3.exceptions import InsecureRequestWarning
import urllib3
urllib3.disable_warnings(InsecureRequestWarning)
 

@frappe.whitelist()
def download():
    filename = 'Form_12'
    build_xlsx_response(filename)

def make_xlsx(sheet_name=None):
    args = frappe.local.form_dict
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name if sheet_name else 'Sheet1'
    
    header_data = get_heading(args)
    for row in header_data:
        ws.append(row)

    data = get_data(args)
    for row in data:
        ws.append(row)
        
        # Inside your loop
        img_url = frappe.db.get_value("File", {'attached_to_name': row[2], "attached_to_doctype": 'Employee'}, ["file_url"])
        if img_url:
            image_url = 'https://103.103.9.11' + img_url
            try:
                response = requests.get(image_url, verify=False)
                if response.status_code == 200 and 'image' in response.headers.get('Content-Type', ''):
                    image_data = response.content
                    img = xlImage(BytesIO(image_data))
                    img.width = 200
                    img.height = 200

                    excel_row = 2 + row[0] 
                    img_cell = ws.cell(row=excel_row, column=18)
                    ws.add_image(img, img_cell.coordinate)
                    ws.row_dimensions[excel_row].height = img.height * 0.75
                else:
                    frappe.log_error(f"Invalid image response: {response.status_code} {response.headers.get('Content-Type')} for {image_url}")
            except Exception as e:
                frappe.log_error(f"Error adding image: {str(e)}")

    merge_cells(ws)
    apply_styles(ws)
    set_column_widths(ws)

    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    xlsx_file.seek(0)
    
    return xlsx_file

def merge_cells(ws):
    ws.merge_cells('A1:F1')
    ws.merge_cells('G1:X1')
    # ws.merge_cells('R1:X1')

def apply_styles(ws):
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    border = Border(
        left=Side(border_style='thin'),
        right=Side(border_style='thin'),
        top=Side(border_style='thin'),
        bottom=Side(border_style='thin')
    )
    header_font = Font(bold=True, size=14)

    for rows in ws.iter_rows(min_row=1, max_row=2, min_col=1, max_col=24):
        for cell in rows:
            cell.font = Font(bold=True)
            cell.alignment = align_center
            cell.border = border
            ws.row_dimensions[cell.row].height = 100

    max_row = ws.max_row
    for rows in ws.iter_rows(min_row=3, min_col=1, max_col=24, max_row=max_row):
        for cell in rows:
            cell.alignment = align_left
            cell.border = border

    

    for rows in ws.iter_rows(min_row=1, max_row=1, min_col=1, max_col=6):
        for cell in rows:
            cell.font = Font(bold=True)
            cell.alignment = align_left
            cell.border = border
    for rows in ws.iter_rows(min_row=1, max_row=1, min_col=7, max_col=24):
        for cell in rows:
            cell.font = header_font
            cell.alignment = align_center
            cell.border = border



def set_column_widths(ws):
    column_widths = [10, 30] + [10] * 20
    for i, width in enumerate(column_widths, start=1):
        ws.column_dimensions[utils.get_column_letter(i)].width = width
    ws.column_dimensions["R"].width = 35
    ws.column_dimensions["G"].width = 30
    ws.column_dimensions["H"].width = 30
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["H"].width = 20
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["F"].width = 20
    ws.column_dimensions["I"].width = 20
    ws.column_dimensions["J"].width = 20
    ws.column_dimensions["K"].width = 20
    ws.column_dimensions["L"].width = 20
    ws.column_dimensions["M"].width = 20
    ws.column_dimensions["N"].width = 20
    ws.column_dimensions["O"].width = 20
    ws.column_dimensions["P"].width = 20
    ws.column_dimensions["Q"].width = 20
    ws.column_dimensions["T"].width = 20
    ws.column_dimensions["S"].width = 20
    ws.column_dimensions["U"].width = 20
    ws.column_dimensions["V"].width = 20
    ws.column_dimensions["W"].width = 20
    ws.column_dimensions["X"].width = 20
    # ws.column_dimensions["F"].width = 20

def build_xlsx_response(filename):
    xlsx_file = make_xlsx(sheet_name=filename)
    frappe.response['filename'] = filename + '.xlsx'
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary'

def get_heading(args):
    company_name = 'Infac India Private Limited'

    header_row = [
        f"Name and Address of the Factory\n{company_name}\n_____\n_____",
        "", "", "", "", "",
        """FORM NO.12\nREGISTER OF ADULT WORKERS AND YOUNG PERSONS\n(Prescribed under rules 80, 86, of the Tamil Nadu Factories Rules, 1950)""",
        "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""
    ]

    header_formatting_row = [
        "S.No", "Name of the worker", "Worker Identity No.", "Gender", "Father/ Spouse Name", "Date of Birth",
        "Present Address", "Permanent Address", "Aadhaar No", "Date of Joining", "Designation", "UAN No",
        "ESIC No", "Date on which completion of 480 days of Services", "Date on which Permanent",
        "Period if suspension if any", "Name, Bank A/C No. Branch , IFSC Code", "Photo", "Mobile No",
        "E-Mail ID", "Specimen Signature / Thumb Impression", "Date of Exit", "Reason for Exit", "Remarks"
    ]

    data_head = [header_row, header_formatting_row]
    return data_head

def get_data(args):
    data = []
    employees = frappe.get_all('Employee', fields=["*"], order_by='employee')

    for i, employee in enumerate(employees, start=1):
        date_of_joining = employee.get('date_of_joining')
        relieving_date = employee.get('relieving_date')
        completion_480_days = ""

        if date_of_joining:
            doj = getdate(date_of_joining)
            target_date = add_days(doj, 480)  
            if relieving_date:
                rel_date = getdate(relieving_date)
                if rel_date >= target_date:
                    completion_480_days = target_date
                else:
                    completion_480_days = '' 
            else:
                completion_480_days = target_date
        completion_480_days = completion_480_days.strftime('%d-%m-%Y') if completion_480_days else ""
        dob = employee.get('date_of_birth', '').strftime('%d-%m-%Y') if employee.get('date_of_birth') else ""
        doj = employee.get('date_of_joining', '').strftime('%d-%m-%Y') if employee.get('date_of_joining') else ""
        relieving_date = employee.get('relieving_date', '').strftime('%d-%m-%Y') if employee.get('relieving_date') else ""
        final_confirmation_date = employee.get('final_confirmation_date', '').strftime('%d-%m-%Y') if employee.get('final_confirmation_date') else ""
        if employee.get('bank_name') and employee.get('bank_ac_no') and employee.get('ifsc_code'):
            bank_details = f"{employee.get('bank_name')}, {employee.get('bank_ac_no')}, {employee.get('ifsc_code')}"
        elif employee.get('bank_name') and employee.get('bank_ac_no'):
            bank_details = f"{employee.get('bank_name')}, {employee.get('bank_ac_no')}"
        elif employee.get('bank_name') and employee.get('ifsc_code'):
            bank_details = f"{employee.get('bank_name')}, {employee.get('ifsc_code')}"
        else:
            bank_details = ""
        row = [
            i, employee.get("employee_name"), employee.get("employee"), employee.get("gender"),
            employee.get("father_name"), dob, employee.get("current_address"), employee.get("permanent_address"),
            employee.get("aadhar_number") or '', doj, employee.get("designation"), employee.get("uan_number") or '', employee.get("esic_no") or "", completion_480_days,final_confirmation_date, "", bank_details, "", employee.get("cell_number") or "", employee.get("user_id") or "", "",relieving_date or "", employee.get("reason_for_leaving") or "", employee.get("feedback") or ""
        ]
        data.append(row)
    
    return data