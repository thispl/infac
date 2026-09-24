# import frappe
# from openpyxl import Workbook
# from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
# from io import BytesIO
# from frappe.utils import getdate

# @frappe.whitelist()
# def download_shift_continue_report(date):
#     if not date:
#         frappe.throw("Date is required")

#     date = getdate(date)

#     # Fetch data, exclude Housekeeping
#     data = frappe.db.sql("""
#         SELECT
#             emp.name AS emp_id,
#             emp.employee_name,
#             emp.department,
#             emp.department_line,
#             att.shift,
#             att.in_time,
#             att.out_time,
#             IFNULL(att.ot_hrs, 0) AS ot_hours,
#             IFNULL(emp.employment_type, '') AS emp_type
#         FROM `tabAttendance` att
#         JOIN `tabEmployee` emp ON emp.name = att.employee
#         WHERE att.attendance_date = %s
#         AND IFNULL(att.ot_hrs, 0) > 0
#         AND emp.department != 'HouseKeeping'
#         ORDER BY emp.department
#     """, (date,), as_dict=True)

#     shifts = ["A", "B", "G", "C"]
#     summary = {}

#     EMPLOYEE_TYPES = {"CONTRACT", "WORKER", "NAPS", "TRAINEE"}

#     # SUMMARY BUILD
#     for row in data:
#         dept = row.department or "Not Set"
#         shift = (row.shift or "G").strip().upper()
#         raw_type = (row.emp_type or "").strip().upper()

#         emp_bucket = "Staff" if raw_type == "STAFF" else "Employee"

#         if dept not in summary:
#             summary[dept] = {
#                 "Employee": {s: 0 for s in shifts},
#                 "Staff": {s: 0 for s in shifts},
#                 "ot_hours": 0
#             }

#         if shift in shifts:
#             summary[dept][emp_bucket][shift] += 1

#         summary[dept]["ot_hours"] += float(row.ot_hours or 0)

#     wb = Workbook()
#     ws1 = wb.active
#     ws1.title = "Shift Continue Report"

#     bold = Font(bold=True)
#     center = Alignment(horizontal="center", vertical="center")
#     thin = Border(
#         left=Side(style="thin"),
#         right=Side(style="thin"),
#         top=Side(style="thin"),
#         bottom=Side(style="thin"),
#     )
#     yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

#     report_date = date.strftime("%d-%m-%Y")
#     ws1.merge_cells("A1:L1")
#     ws1["A1"] = f"Shift Continue Report - {report_date}"
#     ws1["A1"].font = bold
#     ws1["A1"].alignment = center
#     ws1["A1"].fill = yellow_fill

#     # SUMMARY TABLE HEADERS
#     ws1["A2"] = "S.No"
#     ws1["B2"] = "Department"
#     ws1["C2"] = "Employee"
#     ws1["G2"] = "Staff"
#     ws1["K2"] = "Total"
#     ws1["L2"] = "OT Hours"

#     ws1.merge_cells("C2:F2")
#     ws1.merge_cells("G2:J2")
#     ws1.merge_cells("A2:A3")
#     ws1.merge_cells("B2:B3")
#     ws1.merge_cells("K2:K3")
#     ws1.merge_cells("L2:L3")

#     col = 3
#     for s in shifts:
#         ws1.cell(row=3, column=col, value=s)
#         col += 1
#     for s in shifts:
#         ws1.cell(row=3, column=col, value=s)
#         col += 1

#     for row_cells in ws1.iter_rows(min_row=1, max_row=3, min_col=1, max_col=12):
#         for cell in row_cells:
#             cell.font = bold
#             cell.alignment = center
#             cell.border = thin
#             cell.fill = yellow_fill

#     # SUMMARY DATA
#     row_no = 4
#     sno = 1
#     for dept, val in summary.items():
#         emp_total = sum(val["Employee"].values())
#         staff_total = sum(val["Staff"].values())
#         total = emp_total + staff_total

#         ws1.cell(row=row_no, column=1, value=sno)
#         ws1.cell(row=row_no, column=2, value=dept)
#         col = 3
#         for s in shifts:
#             ws1.cell(row=row_no, column=col, value=val["Employee"][s])
#             col += 1
#         for s in shifts:
#             ws1.cell(row=row_no, column=col, value=val["Staff"][s])
#             col += 1
#         ws1.cell(row=row_no, column=11, value=total)
#         ws1.cell(row=row_no, column=12, value=round(val["ot_hours"], 2))

#         for c in range(1, 13):
#             ws1.cell(row=row_no, column=c).border = thin
#             ws1.cell(row=row_no, column=c).alignment = center

#         sno += 1
#         row_no += 1

#     # SHIFT DATA PREP (detailed tables)
#     shift_data = {s: [] for s in shifts}
#     for row in data:
#         s = (row.shift or "G").strip().upper()
#         if s in shift_data:
#             shift_data[s].append(row)

#     # SHIFT TABLE FUNCTION
#     def create_shift_table(shift_name, data_list, s_row, s_col):
#         ws1.cell(row=s_row, column=s_col, value=f"{shift_name} Shift Details").font = bold
#         headers2 = ["S.No", "Employee ID", "Name", "Department Line", "In Time", "Out Time", "OT Hours"]
#         for i, h in enumerate(headers2):
#             cell = ws1.cell(row=s_row + 1, column=s_col + i, value=h)
#             cell.font = bold
#             cell.alignment = center
#             cell.border = thin
#             cell.fill = yellow_fill

#         r = s_row + 2
#         sno = 1
#         for row in data_list:
#             ws1.cell(row=r, column=s_col, value=sno)
#             ws1.cell(row=r, column=s_col + 1, value=row.emp_id)
#             ws1.cell(row=r, column=s_col + 2, value=row.employee_name)
#             ws1.cell(row=r, column=s_col + 3, value=row.department_line)

#             # Set date + HH:MM format
#             ws1.cell(row=r, column=s_col + 4, value=row.in_time).number_format = 'DD-MM-YYYY HH:MM'
#             ws1.cell(row=r, column=s_col + 5, value=row.out_time).number_format = 'DD-MM-YYYY HH:MM'

#             ws1.cell(row=r, column=s_col + 6, value=row.ot_hours)

#             for c in range(s_col, s_col + 7):
#                 ws1.cell(row=r, column=c).border = thin
#                 ws1.cell(row=r, column=c).alignment = center

#             r += 1
#             sno += 1

#     # POSITIONING SHIFT TABLES
#     right_start_col = 14
#     ag_start_row = 1
#     create_shift_table("A", shift_data["A"], ag_start_row, right_start_col)
#     create_shift_table("G", shift_data["G"], ag_start_row, right_start_col + 9)

#     summary_end_row = row_no  
#     b_start_row = summary_end_row + 2
#     create_shift_table("B", shift_data["B"], b_start_row, 1)
#     c_start_row = b_start_row + len(shift_data["B"]) + 4
#     create_shift_table("C", shift_data["C"], c_start_row, 1)

#     # SAVE TO FRAPPE RESPONSE
#     file_stream = BytesIO()
#     wb.save(file_stream)
#     file_stream.seek(0)

#     frappe.response['filename'] = f"Shift_Continue_Report_{report_date}.xlsx"
#     frappe.response['filecontent'] = file_stream.getvalue()
#     frappe.response['type'] = 'binary'




import frappe
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from io import BytesIO
from frappe.utils import getdate

@frappe.whitelist()
def download_shift_continue_report(date):
    if not date:
        frappe.throw("Date is required")

    date = getdate(date)

    # Fetch data, exclude Housekeeping
    data = frappe.db.sql("""
        SELECT
            emp.name AS emp_id,
            emp.employee_name,
            emp.department,
            emp.department_line,
            att.shift,
            att.in_time,
            att.out_time,
            IFNULL(att.ot_hrs, 0) AS ot_hours,
            IFNULL(emp.employment_type, '') AS emp_type
        FROM `tabAttendance` att
        JOIN `tabEmployee` emp ON emp.name = att.employee
        WHERE att.attendance_date = %s
        AND IFNULL(att.ot_hrs, 0) > 0
        AND emp.department != 'HouseKeeping'
        ORDER BY emp.department
    """, (date,), as_dict=True)

    shifts = ["A", "B", "G", "C"]
    summary = {}

    EMPLOYEE_TYPES = {"CONTRACT", "WORKER", "NAPS", "TRAINEE"}

    # SUMMARY BUILD
    for row in data:
        dept = row.department or "Not Set"
        shift = (row.shift or "G").strip().upper()
        raw_type = (row.emp_type or "").strip().upper()

        emp_bucket = "Staff" if raw_type == "STAFF" else "Employee"

        if dept not in summary:
            summary[dept] = {
                "Employee": {s: 0 for s in shifts},
                "Staff": {s: 0 for s in shifts},
                "ot_hours": 0
            }

        if shift in shifts:
            summary[dept][emp_bucket][shift] += 1

        summary[dept]["ot_hours"] += float(row.ot_hours or 0)

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Shift Continue Report"

    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    report_date = date.strftime("%d-%m-%Y")
    ws1.merge_cells("A1:L1")
    ws1["A1"] = f"Shift Continue Report - {report_date}"
    ws1["A1"].font = bold
    ws1["A1"].alignment = center
    ws1["A1"].fill = yellow_fill

    # SUMMARY TABLE HEADERS
    ws1["A2"] = "S.No"
    ws1["B2"] = "Department"
    ws1["C2"] = "Employee/NAPS"
    ws1["G2"] = "Staff"
    ws1["K2"] = "Total"
    ws1["L2"] = "OT Hours"

    ws1.merge_cells("C2:F2")
    ws1.merge_cells("G2:J2")
    ws1.merge_cells("A2:A3")
    ws1.merge_cells("B2:B3")
    ws1.merge_cells("K2:K3")
    ws1.merge_cells("L2:L3")

    col = 3
    for s in shifts:
        ws1.cell(row=3, column=col, value=s)
        col += 1
    for s in shifts:
        ws1.cell(row=3, column=col, value=s)
        col += 1

    for row_cells in ws1.iter_rows(min_row=1, max_row=3, min_col=1, max_col=12):
        for cell in row_cells:
            cell.font = bold
            cell.alignment = center
            cell.border = thin
            cell.fill = yellow_fill

    # SUMMARY DATA
    row_no = 4
    sno = 1
    for dept, val in summary.items():
        emp_total = sum(val["Employee"].values())
        staff_total = sum(val["Staff"].values())
        total = emp_total + staff_total

        ws1.cell(row=row_no, column=1, value=sno)
        ws1.cell(row=row_no, column=2, value=dept)
        col = 3
        for s in shifts:
            ws1.cell(row=row_no, column=col, value=val["Employee"][s])
            col += 1
        for s in shifts:
            ws1.cell(row=row_no, column=col, value=val["Staff"][s])
            col += 1
        ws1.cell(row=row_no, column=11, value=total)
        ws1.cell(row=row_no, column=12, value=round(val["ot_hours"], 2))

        for c in range(1, 13):
            ws1.cell(row=row_no, column=c).border = thin
            ws1.cell(row=row_no, column=c).alignment = center

        sno += 1
        row_no += 1

    total_row = row_no
    last_data_row = row_no - 1
    ws1.cell(row=total_row, column=1, value="TOTAL").font = bold
    ws1.cell(row=total_row, column=1).alignment = center
    ws1.cell(row=total_row, column=1).border = thin
    ws1.cell(row=total_row, column=2).alignment = center
    ws1.cell(row=total_row, column=2).border = thin
    for c in range(3, 13):
        col_letter = get_column_letter(c)
        ws1.cell(row=total_row, column=c, value=f"=SUM({col_letter}4:{col_letter}{last_data_row})")
        ws1.cell(row=total_row, column=c).font = bold
        ws1.cell(row=total_row, column=c).alignment = center
        ws1.cell(row=total_row, column=c).border = thin

    # SHIFT DATA PREP (detailed tables)
    shift_data = {s: [] for s in shifts}
    staff_ot_data = []
    for row in data:
        s = (row.shift or "G").strip().upper()
        raw_type = (row.emp_type or "").strip().upper()
        if raw_type == "STAFF":
            staff_ot_data.append(row)
        elif s in shift_data:
            shift_data[s].append(row)

    # SHIFT TABLE FUNCTION
    def create_ot_table(title, data_list, s_row, s_col):
        ws1.cell(row=s_row, column=s_col, value=title).font = bold
        headers2 = ["S.No", "Employee ID", "Name", "Department Line", "In Time", "Out Time", "OT Hours"]
        for i, h in enumerate(headers2):
            cell = ws1.cell(row=s_row + 1, column=s_col + i, value=h)
            cell.font = bold
            cell.alignment = center
            cell.border = thin
            cell.fill = yellow_fill

        r = s_row + 2
        sno = 1
        total_ot = 0
        for row in data_list:
            ws1.cell(row=r, column=s_col, value=sno)
            ws1.cell(row=r, column=s_col + 1, value=row.emp_id)
            ws1.cell(row=r, column=s_col + 2, value=row.employee_name)
            ws1.cell(row=r, column=s_col + 3, value=row.department_line)

            # Set date + HH:MM format
            ws1.cell(row=r, column=s_col + 4, value=row.in_time).number_format = 'DD-MM-YYYY HH:MM'
            ws1.cell(row=r, column=s_col + 5, value=row.out_time).number_format = 'DD-MM-YYYY HH:MM'

            ws1.cell(row=r, column=s_col + 6, value=row.ot_hours)
            total_ot += float(row.ot_hours or 0)

            for c in range(s_col, s_col + 7):
                ws1.cell(row=r, column=c).border = thin
                ws1.cell(row=r, column=c).alignment = center

            r += 1
            sno += 1

        total_row_local = r
        ws1.cell(row=total_row_local, column=s_col + 5, value="TOTAL").font = bold
        ws1.cell(row=total_row_local, column=s_col + 6, value=total_ot).font = bold
        for c in range(s_col, s_col + 7):
            ws1.cell(row=total_row_local, column=c).border = thin
            ws1.cell(row=total_row_local, column=c).alignment = center

        return total_row_local

    def create_shift_table(shift_name, data_list, s_row, s_col):
        return create_ot_table(f"{shift_name} Shift Details", data_list, s_row, s_col)

    def create_staff_ot_table(data_list, s_row, s_col):
        return create_ot_table("STAFF OT", data_list, s_row, s_col)

    # POSITIONING SHIFT TABLES
    right_start_col = 10
    b_start_row = total_row + 2
    g_start_row = b_start_row

    b_total_row = create_shift_table("B", shift_data["B"], b_start_row, 1)
    g_total_row = create_shift_table("G", shift_data["G"], g_start_row, right_start_col)

    c_start_row = b_total_row + 2
    c_total_row = create_shift_table("C", shift_data["C"], c_start_row, 1)

    a_start_row = c_total_row + 1
    a_total_row = create_shift_table("A", shift_data["A"], a_start_row, 1)

    staff_start_row = g_total_row + 1
    staff_total_row = create_staff_ot_table(staff_ot_data, staff_start_row, right_start_col)

    total_ot_row = a_total_row + 3
    ws1.cell(row=total_ot_row, column=4, value="TOTAL OT").font = bold
    ws1.cell(row=total_ot_row, column=5, value=f"=G{b_total_row}+G{c_total_row}+G{a_total_row}+P{g_total_row}+P{staff_total_row}")
    for c in range(4, 6):
        ws1.cell(row=total_ot_row, column=c).border = thin
        ws1.cell(row=total_ot_row, column=c).alignment = center

    # SAVE TO FRAPPE RESPONSE
    file_stream = BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)

    frappe.response['filename'] = f"Shift_Continue_Report_{report_date}.xlsx"
    frappe.response['filecontent'] = file_stream.getvalue()
    frappe.response['type'] = 'binary'