import frappe
from openpyxl import Workbook, utils
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from io import BytesIO
from frappe.utils import add_days, getdate


@frappe.whitelist()
def download():
    filename = 'Migrant Employee Attendance'
    build_xlsx_response(filename)


def make_xlsx(sheet_name=None):
    args = frappe.local.form_dict

    from_date = getdate(args.get("from_date"))
    to_date = getdate(args.get("to_date"))

    total_days = (to_date - from_date).days + 1
    date_list = [add_days(from_date, i) for i in range(total_days)]

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name if sheet_name else "Sheet1"

    header_data = get_heading(date_list)
    for row in header_data:
        ws.append(row)

    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=8 + len(date_list)
    )

    data = get_data(from_date, to_date, date_list)

    red_fill = PatternFill(start_color="FF9999",
                           end_color="FF9999",
                           fill_type="solid")

    for row in data:
        ws.append(row)

        excel_row = ws.max_row
        start_col = 9  # date columns start

        for col in range(start_col, start_col + len(date_list)):
            cell = ws.cell(row=excel_row, column=col)
            if cell.value == "A":
                cell.fill = red_fill

    apply_styles(ws)
    set_column_widths(ws)

    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    xlsx_file.seek(0)

    return xlsx_file

def get_heading(date_list):

    title_row = ["North Indian Attendance"]

    header_row = [
        "S.No",
        "Emp.Type",
        "Department",
        "Department Line",
        "Emp ID",
        "Employee Name",
        "DOJ",
        "Biometric Pin"
    ]

    header_row.extend([d.strftime("%d-%b") for d in date_list])

    return [title_row, header_row]

def get_data(from_date, to_date, date_list):
    data = []

    employees = frappe.get_all("Employee",filters={ "migrant_labour": "YES", "date_of_joining": ["<=", to_date],"status": ["!=", "Left"] },fields=["name","employee_name","department", "department_line", "employment_type","date_of_joining","attendance_device_id" ],order_by="name" )

    checkins = frappe.get_all( "Employee Checkin",filters={ "log_type": "IN","time": ["between", [from_date, add_days(to_date, 1)]]}, fields=["employee", "time"] )

    checkin_map = {
        (c.employee, getdate(c.time)): True for c in checkins
    }

    

    for i, emp in enumerate(employees, start=1):
        doj =''
        if emp.date_of_joining:
            doj = getdate(emp.date_of_joining).strftime("%d-%m-%Y")
        row = [
            i,
            emp.employment_type,
            emp.department,
            emp.department_line,
            emp.name,
            emp.employee_name,
            doj,
            emp.attendance_device_id
        ]

        emp_doj = getdate(emp.date_of_joining) if emp.date_of_joining else None

        for d in date_list:

            if emp_doj and d < emp_doj:
                row.append("")
                continue

            if checkin_map.get((emp.name, d)):
                row.append("P")
            else:
                row.append("A")

        data.append(row)

    return data

def apply_styles(ws):
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    border = Border(
        left=Side(border_style='thin'),
        right=Side(border_style='thin'),
        top=Side(border_style='thin'),
        bottom=Side(border_style='thin')
    )
    bold = Font(bold=True)
    yellow_fill = PatternFill(
        start_color="FFFF00",
        end_color="FFFF00",
        fill_type="solid"
    )

    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = align_center

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = align_center
        cell.border = border
        cell.fill = yellow_fill
    
    for cell in ws[2]:
        cell.font = bold
            

    for rows in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in rows:
            cell.alignment = align_center
            cell.border = border

def set_column_widths(ws):
    widths = [8, 15, 20, 20, 15, 25, 15, 18]

    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[utils.get_column_letter(i)].width = width

    for col in range(9, ws.max_column + 1):
        ws.column_dimensions[utils.get_column_letter(col)].width = 12

def build_xlsx_response(filename):
    xlsx_file = make_xlsx(sheet_name=filename)
    frappe.response['filename'] = filename + '.xlsx'
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary'