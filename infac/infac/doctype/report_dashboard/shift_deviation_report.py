import frappe
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
import io

@frappe.whitelist()
def download(from_date, to_date):

    if not from_date or not to_date:
        frappe.throw("From Date and To Date are mandatory")

    data = frappe.db.sql("""
        SELECT
            ec.employee,
            emp.employee_name,
            emp.department,
            emp.department_line,
            emp.employee_category,
            ec.time,
            ec.shift,
            ec.assigned_shift
        FROM `tabEmployee Checkin` ec
        LEFT JOIN `tabEmployee` emp ON emp.name = ec.employee
        WHERE DATE(ec.time) BETWEEN %s AND %s
        AND ec.log_type = 'IN'
        AND (
            ec.shift != ec.assigned_shift
            OR ec.shift IS NULL
            OR ec.assigned_shift IS NULL
        )
        ORDER BY ec.time ASC
    """, (from_date, to_date), as_dict=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Shift Mismatch Report"

    headers = [
        "S.No",
        "Emp ID",
        "Emp Name",
        "Department",
        "Department Line",
        "Category",
        "Time",
        "Attended Shift",
        "Assigned Shift"
    ]

    ws.append(headers)

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 25
    ws.column_dimensions["F"].width = 20
    ws.column_dimensions["G"].width = 25
    ws.column_dimensions["I"].width = 18

    for idx, row in enumerate(data, start=1):
        excel_row = ws.max_row + 1

        ws.append([
            idx,
            row.get("employee"),
            row.get("employee_name"),
            row.get("department"),
            row.get("department_line"),
            row.get("employee_category"),
            row.get("time"),
            row.get("shift"),
            row.get("assigned_shift"),
        ])

        for col in range(1, 10):
            cell = ws.cell(row=excel_row, column=col)

            if col == 1:
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.alignment = Alignment(horizontal="left")

            cell.border = thin_border

        ws.cell(row=excel_row, column=7).number_format = "DD-MM-YYYY HH:MM:SS"

    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)

    frappe.response["filename"] = f"Shift_Mismatch_Report_{from_date}_to_{to_date}.xlsx"
    frappe.response["filecontent"] = file_stream.getvalue()
    frappe.response["type"] = "binary"