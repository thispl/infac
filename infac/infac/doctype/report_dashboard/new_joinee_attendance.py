import frappe
from openpyxl import Workbook, utils
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from io import BytesIO
from frappe.utils import add_days, getdate


@frappe.whitelist()
def download():
    filename = 'New Joinee Attendance'
    build_xlsx_response(filename)


def make_xlsx(sheet_name=None):
    args = frappe.local.form_dict

    from_date = getdate(args.get("from_date"))
    to_date = getdate(args.get("to_date"))

    total_days = (to_date - from_date).days + 1
    date_list = [add_days(from_date, i) for i in range(total_days)]

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name if sheet_name else "Sheet1"

    header_data = get_heading(date_list)
    for row in header_data:
        ws.append(row)


    data = get_data(from_date, to_date, date_list)

    red_fill = PatternFill(start_color="FF9999",
                           end_color="FF9999",
                           fill_type="solid")

    for row in data:
        ws.append(row)

        excel_row = ws.max_row
        start_col = 9  # date columns start

        for col in range(start_col, start_col + len(date_list)):
            cell = ws.cell(row=excel_row, column=col)
            if cell.value == "A":
                cell.fill = red_fill

    apply_styles(ws)
    set_column_widths(ws)

    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    xlsx_file.seek(0)

    return xlsx_file

def get_heading(date_list):


    header_row = [
        "S.No",
        "Emp.Type",  
        "Vendor",    
        "Emp ID",
        "Employee Name",
        "DOJ",
        "Department",
        "Department Line"
    ]

    header_row.extend([d.strftime("%d-%b") for d in date_list])

    return [header_row]

def get_data(from_date, to_date, date_list):
    data = []

    employees = frappe.get_all("Employee",filters={ "date_of_joining": ["<=", to_date] , "status": "Active"},fields=["name","employee_name","department", "department_line", "employment_type","date_of_joining","naps_name","employee_category", "holiday_list" ],order_by="name" )

    checkins = frappe.get_all( "Employee Checkin",filters={ "log_type": "IN","time": ["between", [from_date, add_days(to_date, 1)]]}, fields=["employee", "time"] )

    checkin_map = {
        (c.employee, getdate(c.time)): True for c in checkins
    }

    emp_hlists = {e.holiday_list for e in employees if getattr(e, "holiday_list", None)}

    holiday_map = {}

    if emp_hlists:
        holidays = frappe.get_all( "Holiday", filters={"parent": ["in", list(emp_hlists)]}, fields=["parent", "holiday_date", "weekly_off"])
        for h in holidays:
            hdate = getdate(h.holiday_date)
            htype = "WW" if h.weekly_off else "HH"

            holiday_map.setdefault(h.parent, {})
            holiday_map[h.parent][hdate] = htype
    

    for i, emp in enumerate(employees, start=1):
        doj =''
        if emp.date_of_joining:
            doj = getdate(emp.date_of_joining).strftime("%d-%m-%Y")

        vendor = ""
        if emp.employment_type == "Contract":
            vendor = emp.employee_category or ""
        elif emp.employment_type == "NAPS":
            vendor = emp.naps_name or ""

        row = [
            i,
            emp.employment_type,
            vendor,
            emp.name,
            emp.employee_name,
            doj,
            emp.department,
            emp.department_line,
        ]

        emp_doj = getdate(emp.date_of_joining) if emp.date_of_joining else None

        for d in date_list:
            if emp_doj and d < emp_doj:
                row.append("")
                continue
            is_present = checkin_map.get((emp.name, d))

            emp_holidays = holiday_map.get(emp.holiday_list, {})
            is_holiday = emp_holidays.get(d)
            if is_present:
                row.append("P")
            elif is_holiday:
                row.append(is_holiday)
            else:
                row.append("A")

        data.append(row)

    return data

def apply_styles(ws):
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    border = Border(
        left=Side(border_style='thin'),
        right=Side(border_style='thin'),
        top=Side(border_style='thin'),
        bottom=Side(border_style='thin')
    )
    
    for rows in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in rows:
            cell.alignment = align_center
            cell.border = border

def set_column_widths(ws):
    default_width = 7
    custom_widths = {
        2: 10,   # Emp.Type
        3: 10,
        4: 9,
        5: 25,   # Name
        6: 15,   #DOJ
        7: 17 ,    #Dept
        8: 14     #Dept line
    }
    for col in range(1, ws.max_column + 1):
        width = custom_widths.get(col, default_width)
        ws.column_dimensions[utils.get_column_letter(col)].width = width

def build_xlsx_response(filename):
    xlsx_file = make_xlsx(sheet_name=filename)
    frappe.response['filename'] = filename + '.xlsx'
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary'