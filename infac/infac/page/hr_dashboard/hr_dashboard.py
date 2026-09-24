# import frappe


# @frappe.whitelist()
# def get_dept_summary(attendance_date: str = None, shift: str = None):
#     if not attendance_date:
#         attendance_date = frappe.utils.today()

#     filters = {"date": attendance_date}
#     shift_condition = ""
#     if shift and shift != "All":
#         shift_condition = "AND sa.shift_type = %(shift)s"
#         filters["shift"] = shift
#     emp_sql = f"""
#         SELECT
#             COALESCE(e.department, 'Not Set') AS department,
#             sa.shift_type AS shift_code,
#             COUNT(DISTINCT sa.employee) AS available
#         FROM `tabShift Assignment` sa
#         INNER JOIN `tabEmployee` e ON e.name = sa.employee
#         WHERE %(date)s BETWEEN sa.start_date AND IFNULL(sa.end_date, %(date)s)
#           AND sa.docstatus = 1
#           AND (e.employment_type != 'STAFF')
#           {shift_condition}
#         GROUP BY e.department, sa.shift_type
#     """
#     emp_rows = {}
#     for d in frappe.db.sql(emp_sql, filters, as_dict=True):
#         dept = d.department
#         shift_code = d.shift_code or "GENERAL"
#         emp_rows.setdefault(dept, {
#             "overall": 0,
#             "1st Shift": 0,
#             "General": 0,
#             "2nd Shift": 0,
#             "3rd Shift": 0
#         })
#         emp_rows[dept]["overall"] += d.available
#         shift_map = {"A": "1st Shift", "G": "General", "B": "2nd Shift", "C": "3rd Shift"}
#         col = shift_map.get(shift_code, "General")
#         emp_rows[dept][col] += d.available

#     present_sql = f"""
#             SELECT
#                 COALESCE(e.department, 'Not Set') AS department,
#                 ec.shift AS shift_code,
#                 COUNT(DISTINCT ec.employee) AS present
#             FROM `tabEmployee Checkin` ec
#             INNER JOIN `tabEmployee` e ON e.name = ec.employee
#             WHERE DATE(ec.time) = %(date)s
#             AND ec.log_type = 'IN'
#             AND ec.shift IS NOT NULL
#             AND ec.shift != ''
#             AND (e.employment_type != 'STAFF')
#             {"AND ec.shift = %(shift)s" if shift and shift != "All" else ""}
#             GROUP BY e.department, ec.shift
#         """

#     present_rows = {}
#     for d in frappe.db.sql(present_sql, filters, as_dict=True):
#         dept = d.department
#         shift_code = d.shift_code or "G"
#         present_rows.setdefault(dept, {
#             "overall": 0,
#             "1st Shift": 0,
#             "General": 0,
#             "2nd Shift": 0,
#             "3rd Shift": 0
#         })
#         present_rows[dept]["overall"] += d.present
#         shift_map = {"A": "1st Shift", "G": "General", "B": "2nd Shift", "C": "3rd Shift"}
#         col = shift_map.get(shift_code, "General")
#         present_rows[dept][col] += d.present

#     regular_departments = frappe.get_all(
#         "Department",
#         filters={"hr_dashboard": 1, "separate": 0},
#         pluck="name"
#     )
#     separate_departments = frappe.get_all(
#         "Department",
#         filters={"hr_dashboard": 1, "separate": 1},
#         pluck="name"
#     )
#     out = []
#     def get_summary_for_departments(dep_list, label=None):
#         sub_total = {
#             "overall": 0,
#             "present": 0,
#             "absent": 0,
#             "1st Shift": {"plan": 0, "present": 0, "absent": 0},
#             "General": {"plan": 0, "present": 0, "absent": 0},
#             "2nd Shift": {"plan": 0, "present": 0, "absent": 0},
#             "3rd Shift": {"plan": 0, "present": 0, "absent": 0},
#         }
#         rows = []
#         for dep in dep_list:
#             avail = emp_rows.get(dep, {
#                 "overall": 0, "1st Shift": 0, "General": 0, "2nd Shift": 0, "3rd Shift": 0
#             })
#             pres = present_rows.get(dep, {
#                 "overall": 0, "1st Shift": 0, "General": 0, "2nd Shift": 0, "3rd Shift": 0
#             })

#             absent = avail["overall"] - pres["overall"]
#             absent_pct = (absent / avail["overall"] * 100) if avail["overall"] else 0.0

#             row = {
#                 "department": dep,
#                 "available": avail["overall"],
#                 "present": pres["overall"],
#                 "absent": absent,
#                 "absent_pct": round(absent_pct, 1),
#                 "1st Shift": {
#                     "plan": avail["1st Shift"],
#                     "present": pres["1st Shift"],
#                     "absent": avail["1st Shift"] - pres["1st Shift"]
#                 },
#                 "General": {
#                     "plan": avail["General"],
#                     "present": pres["General"],
#                     "absent": avail["General"] - pres["General"]
#                 },
#                 "2nd Shift": {
#                     "plan": avail["2nd Shift"],
#                     "present": pres["2nd Shift"],
#                     "absent": avail["2nd Shift"] - pres["2nd Shift"]
#                 },
#                 "3rd Shift": {
#                     "plan": avail["3rd Shift"],
#                     "present": pres["3rd Shift"],
#                     "absent": avail["3rd Shift"] - pres["3rd Shift"]
#                 },
#                 "bg_color": ""
#             }
#             rows.append(row)
#             sub_total["overall"] += row["available"]
#             sub_total["present"] += row["present"]
#             sub_total["absent"] += row["absent"]

#             for s in ["1st Shift", "General", "2nd Shift", "3rd Shift"]:
#                 sub_total[s]["plan"] += row[s]["plan"]
#                 sub_total[s]["present"] += row[s]["present"]
#                 sub_total[s]["absent"] += row[s]["absent"]

#         if label:
#             rows.append({
#                 "department": label,
#                 "available": sub_total["overall"],
#                 "present": sub_total["present"],
#                 "absent": sub_total["absent"],
#                 "absent_pct": round(
#                     (sub_total["absent"] / sub_total["overall"] * 100)
#                     if sub_total["overall"] else 0, 1
#                 ),
#                 "1st Shift": sub_total["1st Shift"],
#                 "General": sub_total["General"],
#                 "2nd Shift": sub_total["2nd Shift"],
#                 "3rd Shift": sub_total["3rd Shift"],
#                 "bg_color": "#CCE5FF"
#             })
#         return rows, sub_total
#     regular_rows, reg_total = get_summary_for_departments(regular_departments, "SUB TOTAL")
#     out.extend(regular_rows)

#     separate_rows, sep_total = get_summary_for_departments(separate_departments, "SUB TOTAL")
#     out.extend(separate_rows)

#     grand_total = {
#                 "overall": reg_total["overall"] + sep_total["overall"],
#                 "present": reg_total["present"] + sep_total["present"],
#                 "absent": reg_total["absent"] + sep_total["absent"],

#                 "1st Shift": {
#                     "plan": reg_total["1st Shift"]["plan"] + sep_total["1st Shift"]["plan"],
#                     "present": reg_total["1st Shift"]["present"] + sep_total["1st Shift"]["present"],
#                     "absent": reg_total["1st Shift"]["absent"] + sep_total["1st Shift"]["absent"],
#                 },
#                 "General": {
#                     "plan": reg_total["General"]["plan"] + sep_total["General"]["plan"],
#                     "present": reg_total["General"]["present"] + sep_total["General"]["present"],
#                     "absent": reg_total["General"]["absent"] + sep_total["General"]["absent"],
#                 },
#                 "2nd Shift": {
#                     "plan": reg_total["2nd Shift"]["plan"] + sep_total["2nd Shift"]["plan"],
#                     "present": reg_total["2nd Shift"]["present"] + sep_total["2nd Shift"]["present"],
#                     "absent": reg_total["2nd Shift"]["absent"] + sep_total["2nd Shift"]["absent"],
#                 },
#                 "3rd Shift": {
#                     "plan": reg_total["3rd Shift"]["plan"] + sep_total["3rd Shift"]["plan"],
#                     "present": reg_total["3rd Shift"]["present"] + sep_total["3rd Shift"]["present"],
#                     "absent": reg_total["3rd Shift"]["absent"] + sep_total["3rd Shift"]["absent"],
#                 },
#             }

#     out.append({
#         "department": "GRAND TOTAL",
#         "available": grand_total["overall"],
#         "present": grand_total["present"],
#         "absent": grand_total["absent"],
#         "absent_pct": round(
#             (grand_total["absent"] / grand_total["overall"] * 100)
#             if grand_total["overall"] else 0, 1
#         ),
#         "1st Shift": grand_total["1st Shift"],
#         "General": grand_total["General"],
#         "2nd Shift": grand_total["2nd Shift"],
#         "3rd Shift": grand_total["3rd Shift"],
#         "bg_color": "#FFF3CD"
#     })

#     return out

# # @frappe.whitelist()
# # def get_dept_summary(attendance_date: str = None, shift: str = None):
# #     if not attendance_date:
# #         attendance_date = frappe.utils.today()

# #     filters = {"date": attendance_date}
# #     shift_condition = ""
# #     if shift and shift != "All":
# #         shift_condition = "AND sa.shift_type = %(shift)s"
# #         filters["shift"] = shift

# #     emp_sql = f"""
# #         SELECT
# #             COALESCE(e.department, 'Not Set') AS department,
# #             COUNT(DISTINCT sa.employee) AS available
# #         FROM `tabShift Assignment` sa
# #         INNER JOIN `tabEmployee` e ON e.name = sa.employee
# #         WHERE %(date)s BETWEEN sa.start_date AND IFNULL(sa.end_date, %(date)s)
# #           AND sa.docstatus = 1
# #           {shift_condition}
# #         GROUP BY COALESCE(e.department, 'Not Set')
# #     """
# #     emp_rows = {d.department: d.available for d in frappe.db.sql(emp_sql, filters, as_dict=True)}

# #     present_sql = f"""
# #         SELECT
# #             COALESCE(e.department, 'Not Set') AS department,
# #             COUNT(DISTINCT a.employee) AS present
# #         FROM `tabAttendance` a
# #         INNER JOIN `tabEmployee` e ON e.name = a.employee
# #         WHERE a.attendance_date = %(date)s
# #           AND a.docstatus IN (0, 1)
# #           AND a.in_time IS NOT NULL
# #           AND a.in_time != ''
# #           {"AND a.shift = %(shift)s" if shift and shift != "All" else ""}
# #         GROUP BY COALESCE(e.department, 'Not Set')
# #     """
# #     present_rows = {d.department: d.present for d in frappe.db.sql(present_sql, filters, as_dict=True)}

# #     regular_departments = frappe.get_all("Department", filters={"hr_dashboard": 1, "separate": 0}, pluck="name")
# #     separate_departments = frappe.get_all("Department", filters={"hr_dashboard": 1, "separate": 1}, pluck="name")

# #     out = []

# #     def get_summary_for_departments(dep_list, label=None):
# #         sub_total_available = sub_total_present = sub_total_absent = 0
# #         rows = []

# #         for dep in dep_list:
# #             available = int(emp_rows.get(dep, 0))
# #             present = int(present_rows.get(dep, 0))
# #             absent = available - present
# #             absent_pct = (absent / available * 100.0) if available else 0.0

# #             rows.append({
# #                 "department": dep,
# #                 "available": available,
# #                 "present": present,
# #                 "absent": absent,
# #                 "absent_pct": round(absent_pct, 1),
# #                 "bg_color": ""  
# #             })

# #             sub_total_available += available
# #             sub_total_present += present
# #             sub_total_absent += absent

# #         if label:
# #             sub_absent_pct = (sub_total_absent / sub_total_available * 100.0) if sub_total_available else 0.0
# #             rows.append({
# #                 "department": f"{label} ",
# #                 "available": sub_total_available,
# #                 "present": sub_total_present,
# #                 "absent": sub_total_absent,
# #                 "absent_pct": round(sub_absent_pct, 1),
# #                 "bg_color": "#CCE5FF" 
# #             })

# #         return rows, sub_total_available, sub_total_present, sub_total_absent

# #     regular_rows, reg_avail, reg_pres, reg_abs = get_summary_for_departments(regular_departments, "SUB TOTAL")
# #     out.extend(regular_rows)

# #     separate_rows, sep_avail, sep_pres, sep_abs = get_summary_for_departments(separate_departments, "SUB TOTAL")
# #     out.extend(separate_rows)

# #     grand_avail = reg_avail + sep_avail
# #     grand_pres = reg_pres + sep_pres
# #     grand_abs = reg_abs + sep_abs
# #     grand_abs_pct = (grand_abs / grand_avail * 100.0) if grand_avail else 0.0

# #     out.append({
# #         "department": "GRAND TOTAL",
# #         "available": grand_avail,
# #         "present": grand_pres,
# #         "absent": grand_abs,
# #         "absent_pct": round(grand_abs_pct, 1),
# #         "bg_color": "#FFF3CD"
# #     })

# #     return out




# # @frappe.whitelist()
# # def get_late_summary(attendance_date=None):

# #     import frappe
# #     from frappe.utils import today

# #     if not attendance_date:
# #         attendance_date = today()

# #     emp_types = ["STAFF", "WORKER", "NAPS", "Contract", "TRAINEE"]
# #     shifts = ["A", "B", "C", "G"]

# #     sql = """
# #         SELECT 
# #             ec.shift,
# #             emp.employment_type,
# #             COUNT(DISTINCT ec.employee) as late_count
# #         FROM `tabEmployee Checkin` ec
# #         LEFT JOIN `tabEmployee` emp
# #             ON emp.name = ec.employee
# #         WHERE DATE(ec.time) = %(date)s
# #         AND ec.late_entry = 1
# #         GROUP BY ec.shift, emp.employment_type
# #     """

# #     rows = frappe.db.sql(sql, {"date": attendance_date}, as_dict=True)

# #     out = []
# #     totals = {et: 0 for et in emp_types}

# #     for shift in shifts:

# #         row_data = {"Description": shift}

# #         for et in emp_types:

# #             count = next(
# #                 (
# #                     r.late_count
# #                     for r in rows
# #                     if r.shift == shift and r.employment_type == et
# #                 ),
# #                 0,
# #             )

# #             row_data[et] = count
# #             totals[et] += count

# #         out.append(row_data)

# #     total_row = {"Description": "Total"}

# #     for et in emp_types:
# #         total_row[et] = totals[et]

# #     out.append(total_row)

# #     return out


# @frappe.whitelist()
# def get_late_summary(from_date=None, to_date=None):

#     import frappe
#     from frappe.utils import today

#     if not from_date:
#         from_date = today()

#     if not to_date:
#         to_date = from_date

#     emp_types = ["STAFF", "WORKER", "NAPS", "Contract", "TRAINEE"]
#     shifts = ["A", "B", "C", "G"]

#     sql = """
#         SELECT 
#             ec.shift,
#             emp.employment_type,
#             COUNT(DISTINCT ec.employee) as late_count
#         FROM `tabEmployee Checkin` ec
#         LEFT JOIN `tabEmployee` emp
#             ON emp.name = ec.employee
#         WHERE DATE(ec.time) BETWEEN %(from_date)s AND %(to_date)s
#         AND ec.late_entry = 1
#         GROUP BY ec.shift, emp.employment_type
#     """

#     rows = frappe.db.sql(
#         sql,
#         {"from_date": from_date, "to_date": to_date},
#         as_dict=True
#     )

#     out = []
#     totals = {et: 0 for et in emp_types}

#     for shift in shifts:

#         row_data = {"description": shift}

#         for et in emp_types:

#             count = next(
#                 (
#                     r.late_count
#                     for r in rows
#                     if r.shift == shift and r.employment_type == et
#                 ),
#                 0,
#             )

#             row_data[et.lower()] = count
#             totals[et] += count

#         out.append(row_data)

#     total_row = {"description": "Total"}

#     for et in emp_types:
#         total_row[et.lower()] = totals[et]

#     out.append(total_row)

#     return {
#         "summary": out,
#         "total": sum(totals.values())
#     }

# @frappe.whitelist()
# def get_leave_summary(attendance_date: str = None):
#     if not attendance_date:
#         attendance_date = frappe.utils.today()

#     approved = frappe.db.count("Leave Application", {
#         "workflow_state": "Approved",
#         "from_date": ["<=", attendance_date],
#         "to_date": [">=", attendance_date],
#         "half_day": 0
#     })

#     unapproved = frappe.db.count("Leave Application", {
#         "workflow_state": ["!=", "Approved"],
#         "from_date": ["<=", attendance_date],
#         "to_date": [">=", attendance_date],
#         "half_day": 0
#     })

#     approved_hd = frappe.db.count("Leave Application", {
#         "workflow_state": "Approved",
#         "half_day": 1,
#         "half_day_date": attendance_date
#     })

#     unapproved_hd = frappe.db.count("Leave Application", {
#         "workflow_state": ["!=", "Approved"],
#         "half_day": 1,
#         "half_day_date": attendance_date
#     })

#     approved = approved + (approved_hd * 0.5)
#     unapproved = unapproved + (unapproved_hd * 0.5)

#     total = approved + unapproved
#     approved_pct = (approved / total * 100) if total else 0
#     unapproved_pct = (unapproved / total * 100) if total else 0

#     return {
#         "date": attendance_date,
#         "approved": approved,
#         "unapproved": unapproved,
#         "approved_pct": round(approved_pct, 1),
#         "unapproved_pct": round(unapproved_pct, 1),
#     }



# @frappe.whitelist()
# def get_headcount_summary(attendance_date=None):
#     if not attendance_date:
#         attendance_date = frappe.utils.today()

#     data = frappe.db.sql("""
#     SELECT
#         e.department,
#         ec.shift AS shift_code,
#         COUNT(DISTINCT ec.employee) AS total,
#         COUNT(DISTINCT ec.employee) AS present,
#         0 AS absent
#     FROM `tabEmployee Checkin` ec
#     INNER JOIN `tabEmployee` e ON e.name = ec.employee
#     WHERE DATE(ec.time) = %s
#       AND ec.log_type = 'IN'
#       AND ec.shift IS NOT NULL
#     GROUP BY e.department, ec.shift
# """, (attendance_date,), as_dict=True)



#     result = {}

#     for row in data:
#         dept = row.department or "Unknown"
#         shift_code = row.shift_code or "UNKNOWN"

#         if dept not in result:
#             result[dept] = {
#                 "overall": {"avail": 0, "present": 0, "absent": 0},
#                 "1st Shift": {"plan": 0, "present": 0, "absent": 0},
#                 "General": {"plan": 0, "present": 0, "absent": 0},
#                 "2nd Shift": {"plan": 0, "present": 0, "absent": 0},
#                 "3rd Shift": {"plan": 0, "present": 0, "absent": 0},
#             }

#         # overall
#         result[dept]["overall"]["avail"] += row.total
#         result[dept]["overall"]["present"] += row.present
#         result[dept]["overall"]["absent"] += row.absent

#         # shift mapping
#         shift_map = {
#             "A": "1st Shift",
#             "G": "General",
#             "B": "2nd Shift",
#             "C": "3rd Shift"
#         }

#         col = shift_map.get(shift_code)
#         if col:
#             result[dept][col]["plan"] += row.total
#             result[dept][col]["present"] += row.present
#             result[dept][col]["absent"] += row.absent

#     return result



# @frappe.whitelist()
# def get_contractor_summary(attendance_date: str = None, shift: str = None):
#     if not attendance_date:
#         attendance_date = frappe.utils.today()

#     filters = {"date": attendance_date}
#     shift_condition = ""
#     if shift and shift != "All":
#         shift_condition = "AND sa.shift_type = %(shift)s"
#         filters["shift"] = shift
#     else:
#         shift_condition = ""

#     # enrolled_sql = f"""
#     #             SELECT 
#     #                 IFNULL(e.employee_category) AS contractor,
#     #                 COUNT(DISTINCT sa.employee) AS enrolled
#     #             FROM `tabShift Assignment` sa
#     #             INNER JOIN `tabEmployee` e ON e.name = sa.employee
#     #             WHERE sa.docstatus = 1
#     #             AND e.employment_type IN ('Contract', 'NAPS')
#     #             AND %(date)s BETWEEN sa.start_date AND IFNULL(sa.end_date, %(date)s)
#     #             {shift_condition}
#     #             GROUP BY contractor
#     #             """

#     enrolled_sql = f"""
#     SELECT 
#         CASE 
#             WHEN e.employment_type = 'NAPS' THEN n.name
#             ELSE ec.name
#         END AS contractor,
#         COUNT(DISTINCT sa.employee) AS enrolled
#     FROM `tabShift Assignment` sa
#     INNER JOIN `tabEmployee` e ON e.name = sa.employee
#     LEFT JOIN `tabEmployee Category` ec 
#         ON ec.name = e.employee_category AND ec.enable = 1
#     LEFT JOIN `tabNAPS` n 
#         ON n.name = e.naps_name AND n.hr_dashboard= 1
#     WHERE sa.docstatus = 1
#       AND e.employment_type IN ('Contract', 'NAPS')
#       AND %(date)s BETWEEN sa.start_date AND IFNULL(sa.end_date, %(date)s)
#       {shift_condition}
#       AND (
#             (e.employment_type='Contract' AND ec.name IS NOT NULL)
#          OR (e.employment_type='NAPS' AND n.name IS NOT NULL)
#       )
#     GROUP BY contractor
# """
#     enrolled_rows = {d.contractor: d.enrolled for d in frappe.db.sql(enrolled_sql, filters, as_dict=True)}

#     # Present Employees (where In Time is set, filtered by shift)
#     # present_sql = f"""
#     #     SELECT 
#     #         IFNULL(e.employee_category) AS contractor,
#     #         COUNT(DISTINCT a.employee) AS present
#     #     FROM `tabAttendance` a
#     #     INNER JOIN `tabEmployee` e ON e.name = a.employee
#     #     WHERE a.attendance_date = %(date)s
#     #       AND a.docstatus IN (0, 1)
#     #       AND e.employment_type IN ('Contract', 'NAPS')
#     #       AND a.in_time IS NOT NULL
#     #       AND a.in_time != ''
#     #       {"AND a.shift = %(shift)s" if shift and shift != "All" else ""}
#     #     GROUP BY contractor
#     # """

#     present_sql = f"""
#                 SELECT 
#                     CASE 
#                         WHEN e.employment_type = 'NAPS' THEN e.naps_name
#                         ELSE e.employee_category
#                     END AS contractor,
#                     COUNT(DISTINCT a.employee) AS present
#                 FROM `tabAttendance` a
#                 INNER JOIN `tabEmployee` e ON e.name = a.employee
#                 LEFT JOIN `tabEmployee Category` ec 
#                     ON ec.name = e.employee_category AND ec.enable = 1
#                 LEFT JOIN `tabNAPS` n 
#                     ON n.name = e.naps_name AND n.hr_dashboard= 1
#                 WHERE a.attendance_date = %(date)s
#                 AND a.docstatus IN (0, 1)
#                 AND e.employment_type IN ('Contract', 'NAPS')
#                 AND a.in_time IS NOT NULL
#                 AND a.in_time != ''
#                 {"AND a.shift = %(shift)s" if shift and shift != "All" else ""}
#                 AND (
#     (e.employment_type='Contract' AND ec.name IS NOT NULL)
#  OR (e.employment_type='NAPS' AND n.name IS NOT NULL)
# )
#                 GROUP BY contractor
#             """

#     present_rows = {d.contractor: d.present for d in frappe.db.sql(present_sql, filters, as_dict=True)}

#     # Combine all contractors safely (skip None)
#     contractors = sorted(
#         c for c in set(list(enrolled_rows.keys()) + list(present_rows.keys())) if c
#     )

#     out, total_enrolled, total_present, total_absent = [], 0, 0, 0

#     for contractor in contractors:
#         enrolled = int(enrolled_rows.get(contractor, 0))
#         present = int(present_rows.get(contractor, 0))
#         absent = enrolled - present
#         out.append({
#             "contractor": contractor,
#             "enrolled": enrolled,
#             "present": present,
#             "absent": absent,
#         })
#         total_enrolled += enrolled
#         total_present += present
#         total_absent += absent

#     # Grand Total Row
#     out.append({
#         "contractor": "Total",
#         "enrolled": total_enrolled,
#         "present": total_present,
#         "absent": total_absent,
#     })

#     return out


# @frappe.whitelist()
# def get_weekly_attendance(from_date, to_date):
#     from frappe.utils import add_days, getdate
#     from datetime import date
#     frappe.errprint(from_date)
#     if not from_date:
#         from_date = date.today()
#     if not to_date:
#         to_date = add_days(from_date, 6)
#     start = getdate(from_date)
#     end = getdate(to_date)
#     data = []

#     while start <= end:
#         date_str = start.strftime("%d-%m")
#         plan = frappe.db.count("Shift Assignment", {"start_date": start,"docstatus": 1})
#         actual = frappe.db.count("Attendance", {"attendance_date": start, "in_time": ["not in", ["", None]]})
#         gap = plan - actual if plan else 0
#         percent = round((gap / plan) * 100, 1) if plan else 0

#         data.append({
#             "date": date_str,
#             "plan": plan,
#             "actual": actual,
#             "gap": gap, 
#             "percent": percent
#         })
#         start = add_days(start, 1)

#     return data

# @frappe.whitelist()
# def get_dept_line_summary(attendance_date: str = None, shift: str = None):
#     if not attendance_date:
#         attendance_date = frappe.utils.today()

#     filters = {"date": attendance_date}
#     shift_condition = ""
#     if shift and shift != "All":
#         shift_condition = "AND sa.shift_type = %(shift)s"
#         filters["shift"] = shift
#     # Available Employees (based on Shift Assignment)
#     emp_sql = f"""
#         SELECT
#             COALESCE(e.department_line, 'Not Set') AS department_line,
#             COUNT(DISTINCT sa.employee) AS available
#         FROM `tabShift Assignment` sa
#         INNER JOIN `tabEmployee` e ON e.name = sa.employee
#         WHERE %(date)s BETWEEN sa.start_date AND IFNULL(sa.end_date, %(date)s)
#           AND sa.docstatus = 1
#           {shift_condition}
#         GROUP BY COALESCE(e.department_line, 'Not Set')
#     """
#     emp_rows = {d.department_line: d.available for d in frappe.db.sql(emp_sql, filters, as_dict=True)}
#     # Present Employees (where In Time is set)
#     present_sql = f"""
#         SELECT
#             COALESCE(e.department_line, 'Not Set') AS department_line,
#             COUNT(DISTINCT a.employee) AS present
#         FROM `tabAttendance` a
#         INNER JOIN `tabEmployee` e ON e.name = a.employee
#         WHERE a.attendance_date = %(date)s
#         AND a.docstatus IN (0, 1)
#         AND a.in_time IS NOT NULL
#         AND a.in_time != ''
#         {"AND a.shift = %(shift)s" if shift and shift != "All" else ""}
#         GROUP BY COALESCE(e.department_line, 'Not Set')
#     """
#     present_rows = {d.department_line: d.present for d in frappe.db.sql(present_sql, filters, as_dict=True)}
#     #  Department List (only for dashboard)
#     departments = frappe.get_all("Department Line",filters={"hr_dashboard": 1, "dept_line_table":1}, pluck="name") or ["Not Set"]

#     out, total_available, total_present, total_absent = [], 0, 0, 0

#     for dep in departments:
#         available = int(emp_rows.get(dep, 0))
#         present = int(present_rows.get(dep, 0))
#         absent = (available - present) 
#         absent_pct = (absent / available * 100.0) if available else 0.0

#         out.append({
#             "department_line": dep,
#             "available": available,
#             "present": present,
#             "absent": absent,
#             "absent_pct": round(absent_pct, 1),
#         })

#         total_available += available
#         total_present += present
#         total_absent += absent

#     total_absent_pct = (total_absent / total_available * 100.0) if total_available else 0.0
#     out.append({
#         "department_line": "Total",
#         "available": total_available,
#         "present": total_present,
#         "absent": total_absent,
#         "absent_pct": round(total_absent_pct, 1),
#     })

#     return out

# # @frappe.whitelist()
# # def get_permission_summary(from_date, to_date):

# #     data = frappe.db.sql("""
# #         SELECT
# #             employee_id,
# #             employee_name,
# #             session,
# #             from_time,
# #             to_time
# #         FROM `tabPermission Request`
# #         WHERE permission_date BETWEEN %s AND %s
# #     """, (from_date, to_date), as_dict=True)

# #     return {
# #         # "count": len(data),
# #         "details": data
# #     }

# # @frappe.whitelist()
# # def get_leave_app_summary(from_date, to_date):

# #     data = frappe.db.sql("""
# #         SELECT
# #             employee,
# #             employee_name,
# #             leave_type,
# #             description,
# #             superior,
# #             hod,
# #             from_date,
# #             to_date,
# #             half_day, 
# #             total_leave_days,
# #             workflow_state
# #         FROM `tabLeave Application`
# #         WHERE from_date BETWEEN %s AND %s
# #     """, (from_date, to_date), as_dict=True)

# #     return {
# #         # "count": len(data),
# #         "details": data
# #     }

# @frappe.whitelist()
# def get_leave_app_summary(from_date, to_date):

#     data = frappe.db.sql("""
#         SELECT
#             employee,
#             employee_name,
#             leave_type,
#             description,
#             superior,
#             hod,
#             from_date,
#             to_date,
#             half_day,
#             total_leave_days,
#             workflow_state
#         FROM `tabLeave Application`
#         WHERE from_date BETWEEN %s AND %s
#     """, (from_date, to_date), as_dict=True)

#     for row in data:
#         if row.workflow_state == "Incharge Pending":
#             row["approver"] = row.superior or ""
#         elif row.workflow_state == "Superior Pending":
#             row["approver"] = row.hod or ""
#         elif row.workflow_state == "HOD Pending":
#             row["approver"] = row.hod or ""
#         else:
#             row["approver"] = " "

#     return {
#         "details": data
#     }

# @frappe.whitelist()
# def get_late_count():

#     import frappe
#     from frappe.utils import today

#     count = frappe.db.sql("""
#         SELECT COUNT(DISTINCT employee)
#         FROM `tabEmployee Checkin`
#         WHERE DATE(time)=%s
#         AND late_entry=1
#     """, today())[0][0]

#     return count

# # @frappe.whitelist()
# # def get_late_employees():

# #     import frappe
# #     from frappe.utils import today

# #     data = frappe.db.sql("""
# #         SELECT 
# #             ec.employee,
# #             emp.employee_name,
# #             ec.shift,
# #             ec.time,
# #             emp.employment_type
# #         FROM `tabEmployee Checkin` ec
# #         LEFT JOIN `tabEmployee` emp
# #         ON emp.name = ec.employee
# #         WHERE DATE(ec.time)=%(date)s
# #         AND ec.late_entry=1
# #         ORDER BY ec.time
# #     """,{"date":today()},as_dict=True)

# #     return data

# @frappe.whitelist()
# def get_late_employees(from_date=None,to_date=None):

#     import frappe
#     from frappe.utils import today

#     if not from_date:
#         from_date = today()

#     if not to_date:
#         to_date = from_date

#     data = frappe.db.sql("""
#         SELECT 
#             ec.employee,
#             emp.employee_name,
#             ec.shift,
#             ec.time,
#             emp.employment_type
#         FROM `tabEmployee Checkin` ec
#         LEFT JOIN `tabEmployee` emp
#             ON emp.name = ec.employee
#         WHERE DATE(ec.time) BETWEEN %(from_date)s AND %(to_date)s
#         AND ec.late_entry=1
#         ORDER BY ec.time
#     """,{
#         "from_date":from_date,
#         "to_date":to_date
#     },as_dict=True)

#     return data


# @frappe.whitelist()
# def download_late_excel(from_date=None,to_date=None):

#     import frappe
#     import openpyxl
#     from io import BytesIO
#     from frappe.utils import today

#     if not from_date:
#         from_date = today()

#     if not to_date:
#         to_date = from_date

#     data = frappe.db.sql("""
#         SELECT 
#             ec.employee,
#             emp.employee_name,
#             ec.shift,
#             ec.time,
#             emp.employment_type
#         FROM `tabEmployee Checkin` ec
#         LEFT JOIN `tabEmployee` emp
#             ON emp.name = ec.employee
#         WHERE DATE(ec.time) BETWEEN %(from_date)s AND %(to_date)s
#         AND ec.late_entry = 1
#         ORDER BY ec.time
#     """,{
#         "from_date":from_date,
#         "to_date":to_date
#     },as_dict=True)

#     wb = openpyxl.Workbook()
#     ws = wb.active
#     ws.title = "Late Employees"

#     ws.append(["Employee ID","Employee Name","Shift","Checkin Time","Employment Type"])

#     for d in data:
#         ws.append([
#             d.employee,
#             d.employee_name,
#             d.shift,
#             str(d.time),
#             d.employment_type
#         ])

#     file_stream = BytesIO()
#     wb.save(file_stream)

#     frappe.response["filename"] = "Late_Employees.xlsx"
#     frappe.response["filecontent"] = file_stream.getvalue()
#     frappe.response["type"] = "binary"

# @frappe.whitelist()
# def get_permission_count(from_date=None, to_date=None):

#     from frappe.utils import today

#     if not from_date:
#         from_date = today()
#     if not to_date:
#         to_date = from_date

#     from_date = str(from_date)
#     to_date = str(to_date)

#     return frappe.db.sql("""
#         SELECT COUNT(DISTINCT employee_id)
#         FROM `tabPermission Request`
#         WHERE DATE(permission_date) BETWEEN %(from_date)s AND %(to_date)s
#     """, {
#         "from_date": from_date,
#         "to_date": to_date
#     })[0][0]

# @frappe.whitelist()
# def get_permission_summary(from_date=None, to_date=None):

#     from frappe.utils import today

#     if not from_date:
#         from_date = today()
#     if not to_date:
#         to_date = from_date

#     from_date = str(from_date)
#     to_date = str(to_date)

#     summary = frappe.db.sql("""
#         SELECT
#             pr.shift,
#             SUM(CASE WHEN emp.employment_type='STAFF' THEN 1 ELSE 0 END) as staff,
#             SUM(CASE WHEN emp.employment_type='WORKER' THEN 1 ELSE 0 END) as worker,
#             SUM(CASE WHEN emp.employment_type='NAPS' THEN 1 ELSE 0 END) as naps,
#             SUM(CASE WHEN emp.employment_type='Contract' THEN 1 ELSE 0 END) as contract,
#             SUM(CASE WHEN emp.employment_type='Trainee' THEN 1 ELSE 0 END) as trainee
#         FROM `tabPermission Request` pr
#         LEFT JOIN `tabEmployee` emp
#             ON emp.name = pr.employee_id
#         WHERE DATE(pr.permission_date) BETWEEN %(from_date)s AND %(to_date)s
#         GROUP BY pr.shift
#         ORDER BY pr.shift
#     """, {
#         "from_date": from_date,
#         "to_date": to_date
#     }, as_dict=True)

#     # ✅ Total now matches table (count all records, not distinct)
#     total = frappe.db.sql("""
#         SELECT COUNT(name)
#         FROM `tabPermission Request`
#         WHERE DATE(permission_date) BETWEEN %(from_date)s AND %(to_date)s
#     """, {
#         "from_date": from_date,
#         "to_date": to_date
#     })[0][0]

#     return {
#         "summary": summary,
#         "total": total
#     }

# @frappe.whitelist()
# def get_permission_employees(from_date=None, to_date=None):

#     from frappe.utils import today

#     if not from_date:
#         from_date = today()
#     if not to_date:
#         to_date = from_date

#     from_date = str(from_date)
#     to_date = str(to_date)

#     return frappe.db.sql("""
#             SELECT
#                 pr.employee_id,
#                 pr.employee_name,
#                 pr.shift AS shift,
#                 pr.session,
#                 pr.from_time,
#                 pr.to_time,
#                 pr.workflow_state AS status,
#                 pr.creation AS created_on,
#                 CASE 
#                     WHEN pr.workflow_state = 'Incharge Pending' THEN pr.incharge
#                     WHEN pr.workflow_state = 'Superior Pending' THEN pr.superior
#                     ELSE ''
#                 END AS approver
#             FROM `tabPermission Request` pr
#             LEFT JOIN `tabEmployee` emp
#                 ON emp.name = pr.employee_id
#             WHERE DATE(pr.permission_date) BETWEEN %(from_date)s AND %(to_date)s
#             ORDER BY pr.shift, pr.employee_name
#         """, {
#             "from_date": from_date,
#             "to_date": to_date
#         }, as_dict=True)

# @frappe.whitelist()
# def download_permission_excel(from_date=None, to_date=None):

#     import openpyxl
#     from io import BytesIO

#     data = get_permission_employees(from_date, to_date)

#     wb = openpyxl.Workbook()
#     ws = wb.active
#     ws.title = "Permission Employees"

#     ws.append([
#         "Employee ID",
#         "Employee Name",
#         "Shift",
#         "Session",
#         "From Time",
#         "To Time",
#         "Status",
#         "Approver",
#         "Created On"
#     ])

#     for d in data:
#         ws.append([
#             d.employee_id,
#             d.employee_name,
#             d.shift,
#             d.session,
#             str(d.from_time or ""),
#             str(d.to_time or ""),
#             d.status or "",
#             d.approver or "",
#             str(d.created_on or "")
#         ])

#     file_stream = BytesIO()
#     wb.save(file_stream)

#     frappe.response["filename"] = "Permission_Employees.xlsx"
#     frappe.response["filecontent"] = file_stream.getvalue()
#     frappe.response["type"] = "binary"

import frappe
import openpyxl

from io import BytesIO


# =========================================================
# DEPARTMENT SUMMARY
# =========================================================

@frappe.whitelist()
def get_dept_summary(attendance_date: str = None, shift: str = None):

    if not attendance_date:
        attendance_date = frappe.utils.today()

    filters = {
        "date": attendance_date
    }

    shift_condition = ""

    if shift and shift != "All":
        shift_condition = "AND sa.shift_type = %(shift)s"
        filters["shift"] = shift

    emp_sql = f"""
        SELECT
            COALESCE(e.department, 'Not Set') AS department,
            sa.shift_type AS shift_code,
            COUNT(DISTINCT sa.employee) AS available
        FROM `tabShift Assignment` sa
        INNER JOIN `tabEmployee` e
            ON e.name = sa.employee
        WHERE %(date)s BETWEEN sa.start_date
            AND IFNULL(sa.end_date, %(date)s)
          AND sa.docstatus = 1
          AND UPPER(TRIM(COALESCE(e.employment_type, ''))) != 'STAFF'
          {shift_condition}
        GROUP BY e.department, sa.shift_type
    """

    emp_rows = {}

    for d in frappe.db.sql(
        emp_sql,
        filters,
        as_dict=True
    ):

        dept = d.department
        shift_code = (d.shift_code or "G").upper()

        emp_rows.setdefault(
            dept,
            {
                "overall": 0,
                "1st Shift": 0,
                "General": 0,
                "2nd Shift": 0,
                "3rd Shift": 0
            }
        )

        emp_rows[dept]["overall"] += int(d.available or 0)

        shift_map = {
            "A": "1st Shift",
            "G": "General",
            "B": "2nd Shift",
            "C": "3rd Shift"
        }

        col = shift_map.get(
            shift_code,
            "General"
        )

        emp_rows[dept][col] += int(d.available or 0)

    # -----------------------------------------------------
    # PRESENT
    # -----------------------------------------------------

    present_shift_condition = ""

    if shift and shift != "All":
        present_shift_condition = "AND ec.shift = %(shift)s"

    present_sql = f"""
        SELECT
            COALESCE(e.department, 'Not Set') AS department,
            ec.shift AS shift_code,
            COUNT(DISTINCT ec.employee) AS present
        FROM `tabEmployee Checkin` ec
        INNER JOIN `tabEmployee` e
            ON e.name = ec.employee
        WHERE DATE(ec.time) = %(date)s
          AND ec.log_type = 'IN'
          AND ec.shift IS NOT NULL
          AND ec.shift != ''
          AND UPPER(TRIM(COALESCE(e.employment_type, ''))) != 'STAFF'
          {present_shift_condition}
        GROUP BY e.department, ec.shift
    """

    present_rows = {}

    for d in frappe.db.sql(
        present_sql,
        filters,
        as_dict=True
    ):

        dept = d.department
        shift_code = (d.shift_code or "G").upper()

        present_rows.setdefault(
            dept,
            {
                "overall": 0,
                "1st Shift": 0,
                "General": 0,
                "2nd Shift": 0,
                "3rd Shift": 0
            }
        )

        present_rows[dept]["overall"] += int(d.present or 0)

        shift_map = {
            "A": "1st Shift",
            "G": "General",
            "B": "2nd Shift",
            "C": "3rd Shift"
        }

        col = shift_map.get(
            shift_code,
            "General"
        )

        present_rows[dept][col] += int(d.present or 0)

    # -----------------------------------------------------
    # DEPARTMENTS
    # -----------------------------------------------------

    regular_departments = frappe.get_all(
        "Department",
        filters={
            "hr_dashboard": 1,
            "separate": 0
        },
        pluck="name"
    )

    separate_departments = frappe.get_all(
        "Department",
        filters={
            "hr_dashboard": 1,
            "separate": 1
        },
        pluck="name"
    )

    out = []

    # -----------------------------------------------------
    # SUMMARY FUNCTION
    # -----------------------------------------------------

    def get_summary_for_departments(dep_list, label=None):

        sub_total = {
            "overall": 0,
            "present": 0,
            "absent": 0,

            "1st Shift": {
                "plan": 0,
                "present": 0,
                "absent": 0
            },

            "General": {
                "plan": 0,
                "present": 0,
                "absent": 0
            },

            "2nd Shift": {
                "plan": 0,
                "present": 0,
                "absent": 0
            },

            "3rd Shift": {
                "plan": 0,
                "present": 0,
                "absent": 0
            }
        }

        rows = []

        for dep in dep_list:

            avail = emp_rows.get(
                dep,
                {
                    "overall": 0,
                    "1st Shift": 0,
                    "General": 0,
                    "2nd Shift": 0,
                    "3rd Shift": 0
                }
            )

            pres = present_rows.get(
                dep,
                {
                    "overall": 0,
                    "1st Shift": 0,
                    "General": 0,
                    "2nd Shift": 0,
                    "3rd Shift": 0
                }
            )

            absent = (
                avail["overall"]
                - pres["overall"]
            )

            absent_pct = (
                absent
                / avail["overall"]
                * 100
                if avail["overall"]
                else 0.0
            )

            row = {
                "department": dep,
                "available": avail["overall"],
                "present": pres["overall"],
                "absent": absent,

                "absent_pct": round(
                    absent_pct,
                    1
                ),

                "1st Shift": {
                    "plan": avail["1st Shift"],
                    "present": pres["1st Shift"],
                    "absent":
                        avail["1st Shift"]
                        - pres["1st Shift"]
                },

                "General": {
                    "plan": avail["General"],
                    "present": pres["General"],
                    "absent":
                        avail["General"]
                        - pres["General"]
                },

                "2nd Shift": {
                    "plan": avail["2nd Shift"],
                    "present": pres["2nd Shift"],
                    "absent":
                        avail["2nd Shift"]
                        - pres["2nd Shift"]
                },

                "3rd Shift": {
                    "plan": avail["3rd Shift"],
                    "present": pres["3rd Shift"],
                    "absent":
                        avail["3rd Shift"]
                        - pres["3rd Shift"]
                },

                "bg_color": ""
            }

            rows.append(row)

            sub_total["overall"] += row["available"]
            sub_total["present"] += row["present"]
            sub_total["absent"] += row["absent"]

            for s in [
                "1st Shift",
                "General",
                "2nd Shift",
                "3rd Shift"
            ]:

                sub_total[s]["plan"] += row[s]["plan"]
                sub_total[s]["present"] += row[s]["present"]
                sub_total[s]["absent"] += row[s]["absent"]

        if label:

            rows.append(
                {
                    "department": label,

                    "available":
                        sub_total["overall"],

                    "present":
                        sub_total["present"],

                    "absent":
                        sub_total["absent"],

                    "absent_pct": round(
                        (
                            sub_total["absent"]
                            / sub_total["overall"]
                            * 100
                        )
                        if sub_total["overall"]
                        else 0,
                        1
                    ),

                    "1st Shift":
                        sub_total["1st Shift"],

                    "General":
                        sub_total["General"],

                    "2nd Shift":
                        sub_total["2nd Shift"],

                    "3rd Shift":
                        sub_total["3rd Shift"],

                    "bg_color": "#CCE5FF"
                }
            )

        return rows, sub_total

    regular_rows, reg_total = (
        get_summary_for_departments(
            regular_departments,
            "SUB TOTAL"
        )
    )

    out.extend(regular_rows)

    separate_rows, sep_total = (
        get_summary_for_departments(
            separate_departments,
            "SUB TOTAL"
        )
    )

    out.extend(separate_rows)

    # -----------------------------------------------------
    # GRAND TOTAL
    # -----------------------------------------------------

    grand_total = {

        "overall":
            reg_total["overall"]
            + sep_total["overall"],

        "present":
            reg_total["present"]
            + sep_total["present"],

        "absent":
            reg_total["absent"]
            + sep_total["absent"],

        "1st Shift": {
            "plan":
                reg_total["1st Shift"]["plan"]
                + sep_total["1st Shift"]["plan"],

            "present":
                reg_total["1st Shift"]["present"]
                + sep_total["1st Shift"]["present"],

            "absent":
                reg_total["1st Shift"]["absent"]
                + sep_total["1st Shift"]["absent"]
        },

        "General": {
            "plan":
                reg_total["General"]["plan"]
                + sep_total["General"]["plan"],

            "present":
                reg_total["General"]["present"]
                + sep_total["General"]["present"],

            "absent":
                reg_total["General"]["absent"]
                + sep_total["General"]["absent"]
        },

        "2nd Shift": {
            "plan":
                reg_total["2nd Shift"]["plan"]
                + sep_total["2nd Shift"]["plan"],

            "present":
                reg_total["2nd Shift"]["present"]
                + sep_total["2nd Shift"]["present"],

            "absent":
                reg_total["2nd Shift"]["absent"]
                + sep_total["2nd Shift"]["absent"]
        },

        "3rd Shift": {
            "plan":
                reg_total["3rd Shift"]["plan"]
                + sep_total["3rd Shift"]["plan"],

            "present":
                reg_total["3rd Shift"]["present"]
                + sep_total["3rd Shift"]["present"],

            "absent":
                reg_total["3rd Shift"]["absent"]
                + sep_total["3rd Shift"]["absent"]
        }
    }

    out.append(
        {
            "department": "GRAND TOTAL",

            "available":
                grand_total["overall"],

            "present":
                grand_total["present"],

            "absent":
                grand_total["absent"],

            "absent_pct": round(
                (
                    grand_total["absent"]
                    / grand_total["overall"]
                    * 100
                )
                if grand_total["overall"]
                else 0,
                1
            ),

            "1st Shift":
                grand_total["1st Shift"],

            "General":
                grand_total["General"],

            "2nd Shift":
                grand_total["2nd Shift"],

            "3rd Shift":
                grand_total["3rd Shift"],

            "bg_color": "#FFF3CD"
        }
    )

    return out


# =========================================================
# LATE ENTRY SUMMARY
# =========================================================

@frappe.whitelist()
def get_late_summary(
    from_date=None,
    to_date=None
):

    if not from_date:
        from_date = frappe.utils.today()

    if not to_date:
        to_date = from_date

    emp_types = [
        "STAFF",
        "WORKER",
        "NAPS",
        "Contract",
        "Trainee"
    ]

    shifts = [
        "A",
        "B",
        "C",
        "G"
    ]

    sql = """
        SELECT
            UPPER(TRIM(COALESCE(ec.shift, ''))) AS shift,
            emp.employment_type,
            COUNT(DISTINCT ec.employee) AS late_count

        FROM `tabEmployee Checkin` ec

        LEFT JOIN `tabEmployee` emp
            ON emp.name = ec.employee

        WHERE DATE(ec.time)
            BETWEEN %(from_date)s AND %(to_date)s

          AND ec.late_entry = 1

        GROUP BY
            ec.shift,
            emp.employment_type
    """

    rows = frappe.db.sql(
        sql,
        {
            "from_date": from_date,
            "to_date": to_date
        },
        as_dict=True
    )

    out = []

    totals = {
        et: 0
        for et in emp_types
    }

    for shift in shifts:

        row_data = {
            "description": shift
        }

        for et in emp_types:

            count = next(
                (
                    int(r.late_count or 0)
                    for r in rows
                    if r.shift == shift
                    and str(
                        r.employment_type or ""
                    ).strip().upper()
                    == et.upper()
                ),
                0
            )

            row_data[et.lower()] = count
            totals[et] += count

        out.append(row_data)

    total_row = {
        "description": "Total"
    }

    for et in emp_types:
        total_row[et.lower()] = totals[et]

    out.append(total_row)

    return {
        "summary": out,
        "total": sum(totals.values())
    }


# =========================================================
# LEAVE SUMMARY
# =========================================================

@frappe.whitelist()
def get_leave_summary(
    attendance_date: str = None
):

    if not attendance_date:
        attendance_date = frappe.utils.today()

    approved = frappe.db.count(
        "Leave Application",
        {
            "workflow_state": "Approved",
            "from_date": ["<=", attendance_date],
            "to_date": [">=", attendance_date],
            "half_day": 0
        }
    )

    unapproved = frappe.db.count(
        "Leave Application",
        {
            "workflow_state": ["!=", "Approved"],
            "from_date": ["<=", attendance_date],
            "to_date": [">=", attendance_date],
            "half_day": 0
        }
    )

    approved_hd = frappe.db.count(
        "Leave Application",
        {
            "workflow_state": "Approved",
            "half_day": 1,
            "half_day_date": attendance_date
        }
    )

    unapproved_hd = frappe.db.count(
        "Leave Application",
        {
            "workflow_state": ["!=", "Approved"],
            "half_day": 1,
            "half_day_date": attendance_date
        }
    )

    approved = approved + (approved_hd * 0.5)
    unapproved = unapproved + (unapproved_hd * 0.5)

    total = approved + unapproved

    approved_pct = (
        approved / total * 100
        if total
        else 0
    )

    unapproved_pct = (
        unapproved / total * 100
        if total
        else 0
    )

    return {
        "date": attendance_date,
        "approved": approved,
        "unapproved": unapproved,
        "approved_pct": round(
            approved_pct,
            1
        ),
        "unapproved_pct": round(
            unapproved_pct,
            1
        )
    }


# =========================================================
# HEADCOUNT SUMMARY
# =========================================================

@frappe.whitelist()
def get_headcount_summary(
    attendance_date=None
):

    if not attendance_date:
        attendance_date = frappe.utils.today()

    data = frappe.db.sql(
        """
        SELECT
            e.department,
            ec.shift AS shift_code,
            COUNT(DISTINCT ec.employee) AS total,
            COUNT(DISTINCT ec.employee) AS present,
            0 AS absent

        FROM `tabEmployee Checkin` ec

        INNER JOIN `tabEmployee` e
            ON e.name = ec.employee

        WHERE DATE(ec.time) = %s
          AND ec.log_type = 'IN'
          AND ec.shift IS NOT NULL

        GROUP BY
            e.department,
            ec.shift
        """,
        (attendance_date,),
        as_dict=True
    )

    result = {}

    for row in data:

        dept = row.department or "Unknown"

        shift_code = (
            row.shift_code or "UNKNOWN"
        ).upper()

        if dept not in result:

            result[dept] = {

                "overall": {
                    "avail": 0,
                    "present": 0,
                    "absent": 0
                },

                "1st Shift": {
                    "plan": 0,
                    "present": 0,
                    "absent": 0
                },

                "General": {
                    "plan": 0,
                    "present": 0,
                    "absent": 0
                },

                "2nd Shift": {
                    "plan": 0,
                    "present": 0,
                    "absent": 0
                },

                "3rd Shift": {
                    "plan": 0,
                    "present": 0,
                    "absent": 0
                }
            }

        result[dept]["overall"]["avail"] += row.total
        result[dept]["overall"]["present"] += row.present
        result[dept]["overall"]["absent"] += row.absent

        shift_map = {
            "A": "1st Shift",
            "G": "General",
            "B": "2nd Shift",
            "C": "3rd Shift"
        }

        col = shift_map.get(shift_code)

        if col:

            result[dept][col]["plan"] += row.total
            result[dept][col]["present"] += row.present
            result[dept][col]["absent"] += row.absent

    return result


# =========================================================
# CONTRACTOR SUMMARY
# =========================================================

@frappe.whitelist()
def get_contractor_summary(
    attendance_date: str = None,
    shift: str = None
):

    if not attendance_date:
        attendance_date = frappe.utils.today()

    filters = {
        "date": attendance_date
    }

    shift_condition = ""

    if shift and shift != "All":

        shift_condition = """
            AND sa.shift_type = %(shift)s
        """

        filters["shift"] = shift

    enrolled_sql = f"""
        SELECT

            CASE
                WHEN e.employment_type = 'NAPS'
                THEN n.name
                ELSE ec.name
            END AS contractor,

            COUNT(DISTINCT sa.employee) AS enrolled

        FROM `tabShift Assignment` sa

        INNER JOIN `tabEmployee` e
            ON e.name = sa.employee

        LEFT JOIN `tabEmployee Category` ec
            ON ec.name = e.employee_category
            AND ec.enable = 1

        LEFT JOIN `tabNAPS` n
            ON n.name = e.naps_name
            AND n.hr_dashboard = 1

        WHERE sa.docstatus = 1

          AND e.employment_type IN (
              'Contract',
              'NAPS'
          )

          AND %(date)s BETWEEN sa.start_date
              AND IFNULL(sa.end_date, %(date)s)

          {shift_condition}

          AND (
                (
                    e.employment_type = 'Contract'
                    AND ec.name IS NOT NULL
                )
                OR
                (
                    e.employment_type = 'NAPS'
                    AND n.name IS NOT NULL
                )
          )

        GROUP BY contractor
    """

    enrolled_rows = {
        d.contractor: d.enrolled
        for d in frappe.db.sql(
            enrolled_sql,
            filters,
            as_dict=True
        )
    }

    present_sql = f"""
        SELECT

            CASE
                WHEN e.employment_type = 'NAPS'
                THEN e.naps_name
                ELSE e.employee_category
            END AS contractor,

            COUNT(DISTINCT a.employee) AS present

        FROM `tabAttendance` a

        INNER JOIN `tabEmployee` e
            ON e.name = a.employee

        LEFT JOIN `tabEmployee Category` ec
            ON ec.name = e.employee_category
            AND ec.enable = 1

        LEFT JOIN `tabNAPS` n
            ON n.name = e.naps_name
            AND n.hr_dashboard = 1

        WHERE a.attendance_date = %(date)s

          AND a.docstatus IN (0, 1)

          AND e.employment_type IN (
              'Contract',
              'NAPS'
          )

          AND a.in_time IS NOT NULL
          AND a.in_time != ''

          {
              "AND a.shift = %(shift)s"
              if shift and shift != "All"
              else ""
          }

          AND (
                (
                    e.employment_type = 'Contract'
                    AND ec.name IS NOT NULL
                )
                OR
                (
                    e.employment_type = 'NAPS'
                    AND n.name IS NOT NULL
                )
          )

        GROUP BY contractor
    """

    present_rows = {
        d.contractor: d.present
        for d in frappe.db.sql(
            present_sql,
            filters,
            as_dict=True
        )
    }

    contractors = sorted(
        c
        for c in set(
            list(enrolled_rows.keys())
            + list(present_rows.keys())
        )
        if c
    )

    out = []

    total_enrolled = 0
    total_present = 0
    total_absent = 0

    for contractor in contractors:

        enrolled = int(
            enrolled_rows.get(contractor, 0)
        )

        present = int(
            present_rows.get(contractor, 0)
        )

        absent = max(
            enrolled - present,
            0
        )

        out.append(
            {
                "contractor": contractor,
                "enrolled": enrolled,
                "present": present,
                "absent": absent
            }
        )

        total_enrolled += enrolled
        total_present += present
        total_absent += absent

    out.append(
        {
            "contractor": "Total",
            "enrolled": total_enrolled,
            "present": total_present,
            "absent": total_absent
        }
    )

    return out


# =========================================================
# WEEKLY ATTENDANCE
# =========================================================

@frappe.whitelist()
def get_weekly_attendance(
    from_date=None,
    to_date=None
):

    from frappe.utils import add_days, getdate
    from datetime import date

    if not from_date:
        from_date = date.today()

    if not to_date:
        to_date = add_days(from_date, 6)

    start = getdate(from_date)
    end = getdate(to_date)

    data = []

    while start <= end:

        date_str = start.strftime("%d-%m")

        plan = frappe.db.count(
            "Shift Assignment",
            {
                "start_date": start,
                "docstatus": 1
            }
        )

        actual = frappe.db.count(
            "Attendance",
            {
                "attendance_date": start,
                "in_time": [
                    "not in",
                    ["", None]
                ]
            }
        )

        gap = plan - actual if plan else 0

        percent = (
            round(
                (gap / plan) * 100,
                1
            )
            if plan
            else 0
        )

        data.append(
            {
                "date": date_str,
                "plan": plan,
                "actual": actual,
                "gap": gap,
                "percent": percent
            }
        )

        start = add_days(start, 1)

    return data


# =========================================================
# DEPARTMENT LINE SUMMARY
# =========================================================

@frappe.whitelist()
def get_dept_line_summary(
    attendance_date: str = None,
    shift: str = None
):

    if not attendance_date:
        attendance_date = frappe.utils.today()

    filters = {
        "date": attendance_date
    }

    shift_condition = ""

    if shift and shift != "All":

        shift_condition = """
            AND sa.shift_type = %(shift)s
        """

        filters["shift"] = shift

    emp_sql = f"""
        SELECT

            COALESCE(
                e.department_line,
                'Not Set'
            ) AS department_line,

            COUNT(DISTINCT sa.employee) AS available

        FROM `tabShift Assignment` sa

        INNER JOIN `tabEmployee` e
            ON e.name = sa.employee

        WHERE %(date)s BETWEEN sa.start_date
            AND IFNULL(sa.end_date, %(date)s)

          AND sa.docstatus = 1

          {shift_condition}

        GROUP BY COALESCE(
            e.department_line,
            'Not Set'
        )
    """

    emp_rows = {
        d.department_line: d.available
        for d in frappe.db.sql(
            emp_sql,
            filters,
            as_dict=True
        )
    }

    present_sql = f"""
        SELECT

            COALESCE(
                e.department_line,
                'Not Set'
            ) AS department_line,

            COUNT(DISTINCT a.employee) AS present

        FROM `tabAttendance` a

        INNER JOIN `tabEmployee` e
            ON e.name = a.employee

        WHERE a.attendance_date = %(date)s

          AND a.docstatus IN (0, 1)

          AND a.in_time IS NOT NULL
          AND a.in_time != ''

          {
              "AND a.shift = %(shift)s"
              if shift and shift != "All"
              else ""
          }

        GROUP BY COALESCE(
            e.department_line,
            'Not Set'
        )
    """

    present_rows = {
        d.department_line: d.present
        for d in frappe.db.sql(
            present_sql,
            filters,
            as_dict=True
        )
    }

    departments = frappe.get_all(
        "Department Line",
        filters={
            "hr_dashboard": 1,
            "dept_line_table": 1
        },
        pluck="name"
    ) or ["Not Set"]

    out = []

    total_available = 0
    total_present = 0
    total_absent = 0

    for dep in departments:

        available = int(
            emp_rows.get(dep, 0)
        )

        present = int(
            present_rows.get(dep, 0)
        )

        absent = max(
            available - present,
            0
        )

        absent_pct = (
            absent / available * 100.0
            if available
            else 0.0
        )

        out.append(
            {
                "department_line": dep,
                "available": available,
                "present": present,
                "absent": absent,
                "absent_pct": round(
                    absent_pct,
                    1
                )
            }
        )

        total_available += available
        total_present += present
        total_absent += absent

    total_absent_pct = (
        total_absent
        / total_available
        * 100.0
        if total_available
        else 0.0
    )

    out.append(
        {
            "department_line": "Total",
            "available": total_available,
            "present": total_present,
            "absent": total_absent,
            "absent_pct": round(
                total_absent_pct,
                1
            )
        }
    )

    return out


# =========================================================
# LEAVE APPLICATION SUMMARY
# =========================================================

@frappe.whitelist()
def get_leave_app_summary(
    from_date,
    to_date
):

    data = frappe.db.sql(
        """
        SELECT
            employee,
            employee_name,
            leave_type,
            description,
            superior,
            hod,
            from_date,
            to_date,
            half_day,
            total_leave_days,
            workflow_state

        FROM `tabLeave Application`

        WHERE from_date BETWEEN %s AND %s

        ORDER BY from_date
        """,
        (
            from_date,
            to_date
        ),
        as_dict=True
    )

    for row in data:

        if row.workflow_state == "Incharge Pending":

            row["approver"] = (
                row.superior or ""
            )

        elif row.workflow_state == "Superior Pending":

            row["approver"] = (
                row.superior or ""
            )

        elif row.workflow_state == "HOD Pending":

            row["approver"] = (
                row.hod or ""
            )

        else:

            row["approver"] = ""

    return {
        "details": data
    }


# =========================================================
# LATE COUNT
# =========================================================

@frappe.whitelist()
def get_late_count():

    from frappe.utils import today

    count = frappe.db.sql(
        """
        SELECT COUNT(DISTINCT employee)

        FROM `tabEmployee Checkin`

        WHERE DATE(time) = %s
          AND late_entry = 1
        """,
        today()
    )[0][0]

    return count


# =========================================================
# LATE EMPLOYEES
# =========================================================

@frappe.whitelist()
def get_late_employees(
    from_date=None,
    to_date=None,
    shift=None,
    employment_type=None
):

    from frappe.utils import today

    if not from_date:
        from_date = today()

    if not to_date:
        to_date = from_date

    conditions = """
        DATE(ec.time)
        BETWEEN %(from_date)s
        AND %(to_date)s

        AND ec.late_entry = 1
    """

    params = {
        "from_date": str(from_date),
        "to_date": str(to_date)
    }

    if shift and shift not in [
        "Total",
        "All",
        ""
    ]:

        conditions += """
            AND UPPER(
                TRIM(
                    COALESCE(
                        ec.shift,
                        ''
                    )
                )
            )
            =
            UPPER(
                TRIM(
                    %(shift)s
                )
            )
        """

        params["shift"] = shift

    if employment_type and employment_type not in [
        "Total",
        "All",
        ""
    ]:

        conditions += """
            AND UPPER(
                TRIM(
                    COALESCE(
                        emp.employment_type,
                        ''
                    )
                )
            )
            =
            UPPER(
                TRIM(
                    %(employment_type)s
                )
            )
        """

        params["employment_type"] = employment_type

    data = frappe.db.sql(
        f"""
        SELECT
            ec.employee,
            emp.employee_name,
            ec.shift,
            ec.time,
            emp.employment_type

        FROM `tabEmployee Checkin` ec

        LEFT JOIN `tabEmployee` emp
            ON emp.name = ec.employee

        WHERE {conditions}

        ORDER BY ec.time
        """,
        params,
        as_dict=True
    )

    return data


# =========================================================
# DOWNLOAD LATE EXCEL
# =========================================================

@frappe.whitelist()
def download_late_excel(
    from_date=None,
    to_date=None
):

    from frappe.utils import today

    if not from_date:
        from_date = today()

    if not to_date:
        to_date = from_date

    data = get_late_employees(
        from_date=from_date,
        to_date=to_date
    )

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Late Employees"

    ws.append(
        [
            "Employee ID",
            "Employee Name",
            "Shift",
            "Checkin Time",
            "Employment Type"
        ]
    )

    for d in data:

        ws.append(
            [
                d.get("employee") or "",
                d.get("employee_name") or "",
                d.get("shift") or "",
                str(d.get("time") or ""),
                d.get("employment_type") or ""
            ]
        )

    file_stream = BytesIO()

    wb.save(file_stream)

    frappe.response["filename"] = (
        "Late_Employees.xlsx"
    )

    frappe.response["filecontent"] = (
        file_stream.getvalue()
    )

    frappe.response["type"] = "binary"


# =========================================================
# PERMISSION COUNT
# =========================================================

@frappe.whitelist()
def get_permission_count(
    from_date=None,
    to_date=None
):

    from frappe.utils import today

    if not from_date:
        from_date = today()

    if not to_date:
        to_date = from_date

    from_date = str(from_date)
    to_date = str(to_date)

    count = frappe.db.sql(
        """
        SELECT COUNT(DISTINCT pr.employee_id)

        FROM `tabPermission Request` pr

        INNER JOIN `tabEmployee` emp
            ON emp.name = pr.employee_id

        WHERE DATE(pr.permission_date)
            BETWEEN %(from_date)s
            AND %(to_date)s

          AND pr.employee_id IS NOT NULL
          AND pr.employee_id != ''
        """,
        {
            "from_date": from_date,
            "to_date": to_date
        }
    )[0][0]

    return count


# =========================================================
# PERMISSION SUMMARY
# =========================================================

@frappe.whitelist()
def get_permission_summary(
    from_date=None,
    to_date=None
):

    from frappe.utils import today

    if not from_date:
        from_date = today()

    if not to_date:
        to_date = from_date

    from_date = str(from_date)
    to_date = str(to_date)

    rows = frappe.db.sql(
        """
        SELECT

            UPPER(
                TRIM(
                    COALESCE(
                        pr.shift,
                        ''
                    )
                )
            ) AS shift,

            UPPER(
                TRIM(
                    COALESCE(
                        emp.employment_type,
                        ''
                    )
                )
            ) AS employment_type,

            COUNT(
                DISTINCT pr.employee_id
            ) AS employee_count

        FROM `tabPermission Request` pr

        INNER JOIN `tabEmployee` emp
            ON emp.name = pr.employee_id

        WHERE DATE(pr.permission_date)
            BETWEEN %(from_date)s
            AND %(to_date)s

          AND pr.employee_id IS NOT NULL
          AND pr.employee_id != ''

        GROUP BY
            UPPER(
                TRIM(
                    COALESCE(
                        pr.shift,
                        ''
                    )
                )
            ),
            UPPER(
                TRIM(
                    COALESCE(
                        emp.employment_type,
                        ''
                    )
                )
            )
        """,
        {
            "from_date": from_date,
            "to_date": to_date
        },
        as_dict=True
    )

    # -----------------------------------------------------
    # NORMALIZE SHIFT + EMPLOYMENT TYPE
    # -----------------------------------------------------

    count_map = {}

    for row in rows:

        shift_value = (
            row.get("shift") or ""
        ).strip().upper()

        category_value = (
            row.get("employment_type") or ""
        ).strip().upper()

        # Keep the employment type mapping centralized.
        # This prevents Contract/contract mismatch.
        if category_value == "STAFF":

            category_key = "staff"

        elif category_value == "WORKER":

            category_key = "worker"

        elif category_value == "NAPS":

            category_key = "naps"

        elif category_value == "CONTRACT":

            category_key = "contract"

        elif category_value == "TRAINEE":

            category_key = "trainee"

        else:

            continue

        count_map[
            (
                shift_value,
                category_key
            )
        ] = int(
            row.get("employee_count") or 0
        )

    # -----------------------------------------------------
    # FIXED SHIFT ORDER
    # -----------------------------------------------------

    shifts = [
        "A",
        "G",
        "B",
        "C"
    ]

    categories = [
        "staff",
        "worker",
        "naps",
        "contract",
        "trainee"
    ]

    summary = []

    for shift_value in shifts:

        row = {
            "shift": shift_value,
            "description": shift_value
        }

        for category in categories:

            row[category] = count_map.get(
                (
                    shift_value,
                    category
                ),
                0
            )

        summary.append(row)

    # -----------------------------------------------------
    # TOTAL
    # -----------------------------------------------------

    total_row = {
        "shift": "Total",
        "description": "Total"
    }

    for category in categories:

        total_row[category] = 0

    for row in summary:

        for category in categories:

            total_row[category] += int(
                row.get(category) or 0
            )

    summary.append(total_row)

    total = sum(
        total_row[category]
        for category in categories
    )

    return {
        "summary": summary,
        "total": total
    }


# =========================================================
# PERMISSION EMPLOYEES
# =========================================================

@frappe.whitelist()
def get_permission_employees(
    from_date=None,
    to_date=None,
    shift=None,
    employment_type=None
):

    from frappe.utils import today

    if not from_date:
        from_date = today()

    if not to_date:
        to_date = from_date

    from_date = str(from_date)
    to_date = str(to_date)

    conditions = """
        DATE(pr.permission_date)
        BETWEEN %(from_date)s
        AND %(to_date)s
    """

    params = {
        "from_date": from_date,
        "to_date": to_date
    }

    # -----------------------------------------------------
    # SHIFT FILTER
    # -----------------------------------------------------

    if (
        shift
        and str(shift).strip().lower()
        not in [
            "",
            "all",
            "total"
        ]
    ):

        conditions += """
            AND UPPER(
                TRIM(
                    COALESCE(
                        pr.shift,
                        ''
                    )
                )
            )
            =
            UPPER(
                TRIM(
                    %(shift)s
                )
            )
        """

        params["shift"] = str(
            shift
        ).strip()

    # -----------------------------------------------------
    # EMPLOYMENT TYPE FILTER
    # -----------------------------------------------------

    if (
        employment_type
        and str(employment_type).strip().lower()
        not in [
            "",
            "all",
            "total"
        ]
    ):

        conditions += """
            AND UPPER(
                TRIM(
                    COALESCE(
                        emp.employment_type,
                        ''
                    )
                )
            )
            =
            UPPER(
                TRIM(
                    %(employment_type)s
                )
            )
        """

        params["employment_type"] = str(
            employment_type
        ).strip()

    # -----------------------------------------------------
    # FETCH RECORDS
    # -----------------------------------------------------

    data = frappe.db.sql(
        f"""
        SELECT

            pr.employee_id AS employee_id,

            pr.employee_name AS employee_name,

            pr.shift AS shift,

            emp.employment_type AS employment_type,

            pr.session AS session,

            pr.from_time AS from_time,

            pr.to_time AS to_time,

            pr.workflow_state AS status,

            CASE

                WHEN pr.workflow_state =
                    'Incharge Pending'

                THEN COALESCE(
                    pr.incharge,
                    ''
                )

                WHEN pr.workflow_state =
                    'Superior Pending'

                THEN COALESCE(
                    pr.superior,
                    ''
                )

                WHEN pr.workflow_state =
                    'HOD Pending'

                THEN COALESCE(
                    pr.hod,
                    ''
                )

                ELSE ''

            END AS approver,

            pr.creation AS created_on

        FROM `tabPermission Request` pr

        INNER JOIN `tabEmployee` emp
            ON emp.name = pr.employee_id

        WHERE {conditions}

        ORDER BY
            pr.shift,
            pr.employee_name
        """,
        params,
        as_dict=True
    )

    result = []

    for row in data:

        row["employee_id"] = (
            row.get("employee_id")
            or ""
        )

        row["employee_name"] = (
            row.get("employee_name")
            or ""
        )

        row["shift"] = (
            row.get("shift")
            or ""
        )

        row["employment_type"] = (
            row.get("employment_type")
            or ""
        )

        row["session"] = (
            row.get("session")
            or ""
        )

        row["status"] = (
            row.get("status")
            or ""
        )

        row["approver"] = (
            row.get("approver")
            or ""
        )

        row["created_on"] = (
            row.get("created_on")
            or ""
        )

        result.append(row)

    return result


# =========================================================
# DOWNLOAD PERMISSION EXCEL
# =========================================================

@frappe.whitelist()
def download_permission_excel(
    from_date=None,
    to_date=None,
    shift=None,
    employment_type=None
):

    from frappe.utils import today

    if not from_date:
        from_date = today()

    if not to_date:
        to_date = from_date

    data = get_permission_employees(
        from_date=from_date,
        to_date=to_date,
        shift=shift,
        employment_type=employment_type
    )

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Permission Employees"

    ws.append(
        [
            "Employee ID",
            "Employee Name",
            "Shift",
            "Employment Type",
            "Session",
            "From Time",
            "To Time",
            "Status",
            "Approver",
            "Created On"
        ]
    )

    for d in data:

        ws.append(
            [
                d.get("employee_id") or "",
                d.get("employee_name") or "",
                d.get("shift") or "",
                d.get("employment_type") or "",
                d.get("session") or "",
                str(
                    d.get("from_time") or ""
                ),
                str(
                    d.get("to_time") or ""
                ),
                d.get("status") or "",
                d.get("approver") or "",
                str(
                    d.get("created_on") or ""
                )
            ]
        )

    file_stream = BytesIO()

    wb.save(file_stream)

    file_stream.seek(0)

    frappe.response["filename"] = (
        "Permission_Employees.xlsx"
    )

    frappe.response["filecontent"] = (
        file_stream.getvalue()
    )

    frappe.response["type"] = "binary"


# =========================================================
# TRANSACTION HISTORY (LEAVE / LATE / PERMISSION)
# =========================================================

@frappe.whitelist()
def get_transaction_history(
    transaction_type=None,
    employee_type=None,
    employee_id=None,
    start_date=None,
    end_date=None
):

    from frappe.utils import today

    if not start_date:
        start_date = today()

    if not end_date:
        end_date = start_date

    start_date = str(start_date)
    end_date = str(end_date)

    txn_type = str(
        transaction_type or "All"
    ).strip().lower()

    apply_emp_type = (
        employee_type
        and str(employee_type).strip().lower()
        not in ["", "all", "total"]
    )

    apply_emp_id = (
        employee_id
        and str(employee_id).strip() != ""
    )

    records = []

    # -----------------------------------------------------
    # LEAVE APPLICATIONS
    # -----------------------------------------------------

    if txn_type in ["all", "leave"]:

        conditions = """
            la.from_date
            BETWEEN %(start_date)s
            AND %(end_date)s
        """

        params = {
            "start_date": start_date,
            "end_date": end_date
        }

        if apply_emp_type:

            conditions += """
                AND UPPER(
                    TRIM(
                        COALESCE(
                            emp.employment_type,
                            ''
                        )
                    )
                )
                =
                UPPER(
                    TRIM(
                        %(employee_type)s
                    )
                )
            """

            params["employee_type"] = str(
                employee_type
            ).strip()

        if apply_emp_id:

            conditions += """
                AND la.employee = %(employee_id)s
            """

            params["employee_id"] = str(
                employee_id
            ).strip()

        rows = frappe.db.sql(
            f"""
            SELECT
                la.employee AS employee_id,
                la.employee_name AS employee_name,
                emp.employment_type AS employment_type,
                la.from_date AS txn_date,
                '' AS shift,
                la.leave_type AS leave_type,
                la.total_leave_days AS total_leave_days,
                la.workflow_state AS status

            FROM `tabLeave Application` la

            LEFT JOIN `tabEmployee` emp
                ON emp.name = la.employee

            WHERE {conditions}

            ORDER BY la.from_date
            """,
            params,
            as_dict=True
        )

        for row in rows:

            leave_type = (
                row.get("leave_type") or ""
            )

            leave_days = (
                row.get("total_leave_days") or 0
            )

            row["transaction_type"] = "Leave"

            row["details"] = (
                f"{leave_type} "
                f"({leave_days} day(s))"
            ).strip()

            records.append(row)

    # -----------------------------------------------------
    # LATE ENTRIES
    # -----------------------------------------------------

    if txn_type in ["all", "late"]:

        conditions = """
            DATE(ec.time)
            BETWEEN %(start_date)s
            AND %(end_date)s

            AND ec.late_entry = 1
        """

        params = {
            "start_date": start_date,
            "end_date": end_date
        }

        if apply_emp_type:

            conditions += """
                AND UPPER(
                    TRIM(
                        COALESCE(
                            emp.employment_type,
                            ''
                        )
                    )
                )
                =
                UPPER(
                    TRIM(
                        %(employee_type)s
                    )
                )
            """

            params["employee_type"] = str(
                employee_type
            ).strip()

        if apply_emp_id:

            conditions += """
                AND ec.employee = %(employee_id)s
            """

            params["employee_id"] = str(
                employee_id
            ).strip()

        rows = frappe.db.sql(
            f"""
            SELECT
                ec.employee AS employee_id,
                emp.employee_name AS employee_name,
                emp.employment_type AS employment_type,
                DATE(ec.time) AS txn_date,
                ec.shift AS shift,
                ec.time AS checkin_time,
                'Late Entry' AS status

            FROM `tabEmployee Checkin` ec

            LEFT JOIN `tabEmployee` emp
                ON emp.name = ec.employee

            WHERE {conditions}

            ORDER BY ec.time
            """,
            params,
            as_dict=True
        )

        for row in rows:

            checkin_time = row.get(
                "checkin_time"
            )

            row["transaction_type"] = "Late"

            if checkin_time:

                row["details"] = (
                    "Check-in: "
                    + frappe.utils.get_datetime(
                        checkin_time
                    ).strftime(
                        "%d-%m-%Y %H:%M"
                    )
                )

            else:

                row["details"] = ""

            records.append(row)

    # -----------------------------------------------------
    # PERMISSION REQUESTS
    # -----------------------------------------------------

    if txn_type in ["all", "permission"]:

        conditions = """
            DATE(pr.permission_date)
            BETWEEN %(start_date)s
            AND %(end_date)s

            AND pr.employee_id IS NOT NULL
            AND pr.employee_id != ''
        """

        params = {
            "start_date": start_date,
            "end_date": end_date
        }

        if apply_emp_type:

            conditions += """
                AND UPPER(
                    TRIM(
                        COALESCE(
                            emp.employment_type,
                            ''
                        )
                    )
                )
                =
                UPPER(
                    TRIM(
                        %(employee_type)s
                    )
                )
            """

            params["employee_type"] = str(
                employee_type
            ).strip()

        if apply_emp_id:

            conditions += """
                AND pr.employee_id = %(employee_id)s
            """

            params["employee_id"] = str(
                employee_id
            ).strip()

        rows = frappe.db.sql(
            f"""
            SELECT
                pr.employee_id AS employee_id,
                pr.employee_name AS employee_name,
                emp.employment_type AS employment_type,
                DATE(pr.permission_date) AS txn_date,
                pr.shift AS shift,
                pr.session AS session,
                pr.from_time AS from_time,
                pr.to_time AS to_time,
                pr.workflow_state AS status

            FROM `tabPermission Request` pr

            INNER JOIN `tabEmployee` emp
                ON emp.name = pr.employee_id

            WHERE {conditions}

            ORDER BY pr.permission_date
            """,
            params,
            as_dict=True
        )

        for row in rows:

            session = (
                row.get("session") or ""
            )

            from_time = (
                row.get("from_time") or ""
            )

            to_time = (
                row.get("to_time") or ""
            )

            time_range = (
                f"{from_time} - {to_time}"
                if from_time or to_time
                else ""
            )

            row["transaction_type"] = (
                "Permission"
            )

            row["details"] = (
                f"{session} {time_range}"
            ).strip()

            records.append(row)

    # -----------------------------------------------------
    # SORT + COUNT
    # -----------------------------------------------------

    type_order = {
        "Late": 0,
        "Leave": 1,
        "Permission": 2
    }

    records.sort(
        key=lambda r: (
            type_order.get(
                r.get("transaction_type"),
                3
            ),
            str(r.get("txn_date") or ""),
            str(r.get("employee_id") or "")
        )
    )

    return {
        "count": len(records),
        "records": records
    }


# =========================================================
# DOWNLOAD TRANSACTION HISTORY EXCEL
# =========================================================

@frappe.whitelist()
def download_transaction_history(
    transaction_type=None,
    employee_type=None,
    employee_id=None,
    start_date=None,
    end_date=None
):

    result = get_transaction_history(
        transaction_type=transaction_type,
        employee_type=employee_type,
        employee_id=employee_id,
        start_date=start_date,
        end_date=end_date
    )

    data = result.get("records") or []

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Transaction History"

    ws.append(
        [
            "Type",
            "Employee ID",
            "Employee Name",
            "Employee Type",
            "Date",
            "Shift",
            "Details",
            "Status"
        ]
    )

    for d in data:

        txn_date = d.get("txn_date")

        if txn_date:

            txn_date = (
                frappe.utils.getdate(
                    txn_date
                ).strftime(
                    "%d-%m-%Y"
                )
            )

        ws.append(
            [
                d.get("transaction_type") or "",
                d.get("employee_id") or "",
                d.get("employee_name") or "",
                d.get("employment_type") or "",
                txn_date or "",
                d.get("shift") or "",
                d.get("details") or "",
                d.get("status") or ""
            ]
        )

    file_stream = BytesIO()

    wb.save(file_stream)

    file_stream.seek(0)

    frappe.response["filename"] = (
        "Transaction_History.xlsx"
    )

    frappe.response["filecontent"] = (
        file_stream.getvalue()
    )

    frappe.response["type"] = "binary"


# =========================================================
# TRANSACTION CALENDAR (MONTH-WISE)
# =========================================================

@frappe.whitelist()
def get_txn_calendar(
    application=None,
    employee_type=None,
    employee_id=None,
    from_date=None,
    to_date=None
):

    from frappe.utils import (
        today,
        getdate,
        get_datetime,
        add_days
    )

    today_str = today()

    if not from_date:

        from_date = f"{today_str[:4]}-01-01"

    if not to_date:

        to_date = f"{today_str[:4]}-12-31"

    from_date = str(getdate(from_date))
    to_date = str(getdate(to_date))

    if from_date > to_date:

        from_date, to_date = (
            to_date,
            from_date
        )

    year_start = from_date
    year_end = to_date

    app_type = str(
        application or "All"
    ).strip().lower()

    apply_emp_type = (
        employee_type
        and str(employee_type).strip().lower()
        not in ["", "all", "total"]
    )

    emp_type_condition = ""

    params = {
        "year_start": year_start,
        "year_end": year_end
    }

    if apply_emp_type:

        emp_type_condition = """
            AND UPPER(
                TRIM(
                    COALESCE(
                        emp.employment_type,
                        ''
                    )
                )
            )
            =
            UPPER(
                TRIM(
                    %(employee_type)s
                )
            )
        """

        params["employee_type"] = str(
            employee_type
        ).strip()

    apply_emp_id = (
        employee_id
        and str(employee_id).strip() != ""
    )

    emp_id_leave = ""
    emp_id_late = ""
    emp_id_permission = ""

    if apply_emp_id:

        params["employee_id"] = str(
            employee_id
        ).strip()

        emp_id_leave = (
            "AND la.employee = %(employee_id)s"
        )

        emp_id_late = (
            "AND ec.employee = %(employee_id)s"
        )

        emp_id_permission = (
            "AND pr.employee_id = %(employee_id)s"
        )

    employees = {}

    def add_txn(
        employee_id,
        employee_name,
        employment_type,
        date_str,
        code,
        txn_kind,
        details
    ):

        if not employee_id or not date_str:
            return

        date_str = str(date_str)

        if (
            date_str < year_start
            or date_str > year_end
        ):
            return

        entry = employees.setdefault(
            employee_id,
            {
                "employee_id": employee_id,
                "employee_name": (
                    employee_name or ""
                ),
                "employment_type": (
                    employment_type or ""
                ),
                "days": {}
            }
        )

        day_list = entry["days"].setdefault(
            date_str,
            []
        )

        day_list.append(
            {
                "code": code,
                "type": txn_kind,
                "details": details or ""
            }
        )

    # -----------------------------------------------------
    # LATE ENTRIES
    # -----------------------------------------------------

    if app_type in ["all", "late"]:

        rows = frappe.db.sql(
            f"""
            SELECT
                ec.employee AS employee_id,
                emp.employee_name AS employee_name,
                emp.employment_type AS employment_type,
                DATE(ec.time) AS txn_date,
                ec.time AS checkin_time

            FROM `tabEmployee Checkin` ec

            LEFT JOIN `tabEmployee` emp
                ON emp.name = ec.employee

            WHERE DATE(ec.time)
                BETWEEN %(year_start)s
                AND %(year_end)s

              AND ec.late_entry = 1

            {emp_type_condition}

            {emp_id_late}
            """,
            params,
            as_dict=True
        )

        for row in rows:

            checkin_time = row.get(
                "checkin_time"
            )

            details = ""

            if checkin_time:

                details = (
                    "Check-in: "
                    + get_datetime(
                        checkin_time
                    ).strftime(
                        "%d-%m-%Y %H:%M"
                    )
                )

            add_txn(
                row.get("employee_id"),
                row.get("employee_name"),
                row.get("employment_type"),
                row.get("txn_date"),
                "LT",
                "Late",
                details
            )

    # -----------------------------------------------------
    # LEAVE APPLICATIONS
    # -----------------------------------------------------

    if app_type in ["all", "leave"]:

        rows = frappe.db.sql(
            f"""
            SELECT
                la.employee AS employee_id,
                la.employee_name AS employee_name,
                emp.employment_type AS employment_type,
                la.from_date AS from_date,
                la.to_date AS to_date,
                la.leave_type AS leave_type,
                la.total_leave_days AS total_leave_days,
                la.workflow_state AS status

            FROM `tabLeave Application` la

            LEFT JOIN `tabEmployee` emp
                ON emp.name = la.employee

            WHERE la.from_date <= %(year_end)s
              AND la.to_date >= %(year_start)s

            {emp_type_condition}

            {emp_id_leave}
            """,
            params,
            as_dict=True
        )

        for row in rows:

            from_d = getdate(
                row.get("from_date")
            )

            to_d = getdate(
                row.get("to_date") or
                row.get("from_date")
            )

            leave_type = (
                row.get("leave_type") or "Leave"
            )

            status = (
                row.get("status") or ""
            )

            details = (
                f"{leave_type}"
                + (
                    f" - {status}"
                    if status
                    else ""
                )
            )

            cur = from_d

            while cur <= to_d:

                add_txn(
                    row.get("employee_id"),
                    row.get("employee_name"),
                    row.get("employment_type"),
                    str(cur),
                    "L",
                    "Leave",
                    details
                )

                cur = add_days(cur, 1)

    # -----------------------------------------------------
    # PERMISSION REQUESTS
    # -----------------------------------------------------

    if app_type in ["all", "permission"]:

        rows = frappe.db.sql(
            f"""
            SELECT
                pr.employee_id AS employee_id,
                pr.employee_name AS employee_name,
                emp.employment_type AS employment_type,
                DATE(pr.permission_date) AS txn_date,
                pr.session AS session,
                pr.from_time AS from_time,
                pr.to_time AS to_time,
                pr.workflow_state AS status

            FROM `tabPermission Request` pr

            INNER JOIN `tabEmployee` emp
                ON emp.name = pr.employee_id

            WHERE DATE(pr.permission_date)
                BETWEEN %(year_start)s
                AND %(year_end)s

              AND pr.employee_id IS NOT NULL
              AND pr.employee_id != ''

            {emp_type_condition}

            {emp_id_permission}
            """,
            params,
            as_dict=True
        )

        for row in rows:

            session = (
                row.get("session") or ""
            )

            from_time = (
                row.get("from_time") or ""
            )

            to_time = (
                row.get("to_time") or ""
            )

            status = (
                row.get("status") or ""
            )

            time_range = (
                f"{from_time} - {to_time}"
                if from_time or to_time
                else ""
            )

            details = (
                f"{session} {time_range}"
            ).strip()

            if status:

                details = (
                    f"{details} - {status}"
                    if details
                    else status
                )

            add_txn(
                row.get("employee_id"),
                row.get("employee_name"),
                row.get("employment_type"),
                row.get("txn_date"),
                "P",
                "Permission",
                details
            )

    # -----------------------------------------------------
    # SORTED OUTPUT
    # -----------------------------------------------------

    result = sorted(
        employees.values(),
        key=lambda e: str(
            e.get("employee_id") or ""
        )
    )

    return {
        "from_date": year_start,
        "to_date": year_end,
        "employees": result
    }


# =========================================================
# DOWNLOAD TRANSACTION CALENDAR EXCEL
# =========================================================

@frappe.whitelist()
def download_txn_calendar(
    application=None,
    employee_type=None,
    employee_id=None,
    from_date=None,
    to_date=None
):

    from frappe.utils import (
        getdate,
        add_months
    )

    response = get_txn_calendar(
        application=application,
        employee_type=employee_type,
        employee_id=employee_id,
        from_date=from_date,
        to_date=to_date
    )

    from_d = getdate(response.get("from_date"))
    to_d = getdate(response.get("to_date"))

    employees = response.get("employees") or []

    # month keys + labels across the selected range
    month_keys = []
    month_labels = []

    cur = from_d.replace(day=1)

    while cur <= to_d:

        month_keys.append((cur.year, cur.month))
        month_labels.append(
            cur.strftime("%b %Y")
        )

        cur = add_months(cur, 1)

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Calendar"

    ws.append(
        ["Employee ID", "Employee Name"]
        + month_labels
        + ["Total"]
    )

    month_totals = [0] * len(month_keys)
    grand_total = 0

    for emp in employees:

        days = emp.get("days") or {}

        row = [
            emp.get("employee_id") or "",
            emp.get("employee_name") or ""
        ]

        row_total = 0

        for i, (y, m) in enumerate(month_keys):

            prefix = "%04d-%02d-" % (y, m)

            count = sum(
                len(entries)
                for key, entries in days.items()
                if str(key).startswith(prefix)
            )

            row.append(count or "")
            row_total += count
            month_totals[i] += count

        row.append(row_total)
        grand_total += row_total

        ws.append(row)

    ws.append(
        ["", "Total"]
        + [t or "" for t in month_totals]
        + [grand_total]
    )

    file_stream = BytesIO()

    wb.save(file_stream)

    file_stream.seek(0)

    frappe.response["filename"] = (
        "Transaction_Calendar.xlsx"
    )

    frappe.response["filecontent"] = (
        file_stream.getvalue()
    )

    frappe.response["type"] = "binary"